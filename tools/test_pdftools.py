"""Engine test of the PDF tools: forms, pages, bookmarks, links, properties, security,
grayscale, PDF/A, export, digital signatures, repair and batch. Runs anywhere (no Windows
OCR needed).   python tools/test_pdftools.py"""
import io
import os
import shutil
import sys
import traceback

# the Windows console (cp1252) cannot print every character of the results
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
WORK = os.path.join(ROOT, "build", "pdftools")
shutil.rmtree(WORK, ignore_errors=True)
os.makedirs(WORK)
os.environ["CALAMO_DATA"] = os.path.join(WORK, "data")
sys.path.insert(0, os.path.join(ROOT, "app"))
import pymupdf as fitz  # noqa: E402
from engine.document import Session  # noqa: E402
from engine import batch, files, signing  # noqa: E402

failed = 0


def ok(name, cond, extra=""):
    global failed
    print(("PASS " if cond else "FAIL ") + name + (f"  [{extra}]" if extra != "" else ""))
    if not cond:
        failed += 1


def check(name, fn):
    try:
        fn()
    except Exception:
        global failed
        failed += 1
        print("FAIL " + name + " (eccezione)")
        traceback.print_exc()


def P(name):
    return os.path.join(WORK, name)


def make_form():
    d = fitz.open()
    p = d.new_page(width=595, height=842)
    p.insert_text((60, 60), "Modulo di iscrizione", fontsize=18)
    p.insert_text((60, 110), "Nome: ____________________", fontsize=11)
    p.insert_text((60, 140), "Cognome: .......................................", fontsize=11)
    p.insert_text((60, 170), "Città", fontsize=11)
    p.draw_line((110, 172), (300, 172), color=(0, 0, 0), width=0.8)
    p.draw_rect((60, 190, 72, 202), color=(0, 0, 0), width=0.8)
    p.insert_text((78, 200), "Accetto le condizioni", fontsize=11)
    # a table must not become fields
    for i in range(3):
        p.draw_line((60, 300 + i * 20), (400, 300 + i * 20), color=(0, 0, 0))
    for x in (60, 230, 400):
        p.draw_line((x, 300), (x, 340), color=(0, 0, 0))
    p.insert_text((65, 315), "Voce", fontsize=10)
    # some text with a heading for bookmarks
    for k in range(3):
        q = d.new_page(width=595, height=842)
        q.insert_text((60, 80), f"Capitolo {k + 1}", fontsize=20)
        q.insert_text((60, 120), "Testo del capitolo con alcune parole di contenuto normale.", fontsize=11)
        q.insert_text((60, 140), "Altro testo del capitolo per avere un corpo del testo.", fontsize=11)
    d.save(P("form.pdf"))


def t_forms():
    s = Session()
    s.open(P("form.pdf"))
    r = s.detect_fields("1")
    names = sorted(f["name"] for f in r["fields"])
    kinds = sorted(f["kind"] for f in r["fields"])
    ok("riconoscimento campi: 4 trovati", r["added"] == 4, f"{names} {kinds}")
    ok("nomi dai testi accanto", "Nome" in names and "Cognome" in names and "Città" in names, names)
    ok("casella di spunta", "checkbox" in kinds)
    f = {x["name"]: x for x in r["fields"]}
    vals = [{"page": 0, "id": f["Nome"]["id"], "value": "Mario"}]
    cb = next(x for x in r["fields"] if x["kind"] == "checkbox")
    vals.append({"page": 0, "id": cb["id"], "value": True})
    r = s.fill_fields(vals)
    f = {x["name"]: x for x in r["fields"]}
    ok("compila testo", f["Nome"]["value"] == "Mario")
    ok("compila casella", next(x for x in r["fields"] if x["kind"] == "checkbox")["value"] is True)
    r = s.add_field(0, "combo", [300, 400, 450, 420], "Provincia", ["MI", "RM", "TO"])
    pv = next(x for x in r["fields"] if x["name"] == "Provincia")
    ok("nuovo menu a tendina", pv["kind"] == "combo" and pv["options"] == ["MI", "RM", "TO"] and pv["value"] == "MI")
    s.fill_fields([{"page": 0, "id": pv["id"], "value": "RM"}])
    ok("scelta menu", next(x for x in s.fields() if x["name"] == "Provincia")["value"] == "RM")
    s.add_field(0, "radio", [300, 450, 312, 462], "Sesso")
    r2 = s.add_field(0, "radio", [340, 450, 352, 462], "Sesso")
    rad = [x for x in r2["fields"] if x["kind"] == "radio"]
    ok("gruppo di scelte", len(rad) == 2 and rad[0]["name"] == rad[1]["name"], [x["name"] for x in rad])
    s.fill_fields([{"page": 0, "id": rad[1]["id"], "value": True}])
    rad = [x for x in s.fields() if x["kind"] == "radio"]
    ok("scelta singola", [x["value"] for x in rad] == [False, True], [x["value"] for x in rad])
    s.add_field(0, "signature", [300, 600, 500, 660], "Firma del richiedente")
    s.delete_field(0, pv["id"])
    ok("campo eliminato", not any(x["name"] == "Provincia" for x in s.fields()))
    s.save(P("form-compilato.pdf"))
    d = fitz.open(P("form-compilato.pdf"))
    vals = {w.field_name: w.field_value for w in d[0].widgets()}
    ok("valori salvati nel file", vals.get("Nome") == "Mario", vals)
    s.undo()
    ok("annulla", any(x["name"] == "Provincia" for x in s.fields()))
    s.redo()
    s.flatten(True, False)
    ok("appiattisci moduli", not s.fields() and "Mario" in s.doc[0].get_text())


def t_pages():
    s = Session()
    s.open(P("form.pdf"))
    r = s.crop_pages("1", margins=[10, 10, 10, 10])
    ok("ritaglia margini", abs(r["pages"][0]["w"] - (595 - 20 * 72 / 25.4)) < 1, r["pages"][0])
    s.reset_crop("1")
    ok("annulla ritaglio", abs(s.info()["pages"][0]["w"] - 595) < 0.5)
    s.doc[1].set_rotation(90)
    s.crop_pages("2", rect=[50, 40, 500, 400])
    ok("ritaglia rettangolo su pagina ruotata", abs(s.info()["pages"][1]["w"] - 450) < 1 and abs(s.info()["pages"][1]["h"] - 360) < 1,
       s.info()["pages"][1])
    s.crop_pages("2", rect=[10, 10, 300, 200])          # crop of a cropped page
    ok("ritaglio di una pagina già ritagliata", abs(s.info()["pages"][1]["w"] - 290) < 1, s.info()["pages"][1])
    r = s.crop_pages("3", auto=True)
    ok("ritaglio automatico", r["cropped"] == 1 and r["pages"][2]["w"] < 595)
    r = s.resize_pages("4", "A3")
    ok("ridimensiona in A3", abs(r["pages"][3]["w"] - 841.89) < 1 and "Capitolo 3" in s.doc[3].get_text())
    ok("ridimensiona: stesse pagine", len(s.doc) == 4)
    nup = fitz.open(stream=s.nup_bytes(2, 1, "A4"), filetype="pdf")
    ok("2 pagine per foglio", len(nup) == 2 and nup[0].rect.width > nup[0].rect.height, [len(nup), nup[0].rect])
    ok("4 per foglio", len(fitz.open(stream=s.nup_bytes(2, 2, "A4"), filetype="pdf")) == 1)
    bk = fitz.open(stream=s.booklet_bytes("A4"), filetype="pdf")
    t0 = bk[0].get_text()
    ok("opuscolo", len(bk) == 2 and "Capitolo 3" in t0 and "Modulo" in t0, [len(bk), t0[:60]])


def t_bookmarks_links_props():
    s = Session()
    s.open(P("form.pdf"))
    prop = s.bookmarks_from_headings()
    titles = [b["title"] for b in prop]
    ok("titoli trovati", all(f"Capitolo {k}" in titles for k in (1, 2, 3)), titles)
    s.set_bookmarks(prop + [{"level": 5, "title": "Sotto", "page": 99}])
    bm = s.bookmarks()
    ok("segnalibri salvati e corretti", bm[-1]["level"] <= bm[-2]["level"] + 1 and bm[-1]["page"] == 3, bm[-2:])
    r = s.add_link(0, [60, 40, 260, 65], uri="www.esempio.it")
    ok("link web", r["links"][0]["uri"] == "https://www.esempio.it", r["links"])
    r = s.add_link(0, [60, 95, 260, 115], target_page=2)
    ok("link a pagina", any(l["kind"] == "page" and l["page"] == 2 for l in r["links"]), r["links"])
    r = s.delete_link(0, r["links"][0]["id"])
    ok("link eliminato", len(r["links"]) == 1)
    r = s.set_properties({"title": "Iscrizione", "author": "Ufficio", "lang": "it-IT"})
    ok("proprietà", r["properties"]["title"] == "Iscrizione" and r["properties"]["lang"] == "it-IT")
    r = s.stamp_pages({"footer": "{bates}", "align": "right", "bates_prefix": "CAUSA-", "bates_digits": 5, "bates_start": 7})
    ok("numerazione Bates", "CAUSA-00007" in s.doc[0].get_text() and "CAUSA-00010" in s.doc[3].get_text())
    s.save(P("segnalibri.pdf"))
    d = fitz.open(P("segnalibri.pdf"))
    ok("segnalibri e link nel file", len(d.get_toc()) == 5 and d[0].first_link is not None)


def t_security_color():
    s = Session()
    s.open(P("form.pdf"))
    s.save_protected(P("protetto.pdf"), "segreta", "", True, False, False)
    s2 = Session()
    try:
        s2.open(P("protetto.pdf"))
        ok("serve la password", False)
    except RuntimeError as e:
        ok("serve la password", str(e) == "PASSWORD_REQUIRED")
    s2.open(P("protetto.pdf"), "segreta")
    s2.add_text(0, 300, 500, "Aggiunto", {"family": "Arial", "size": 12, "color": "#000000"})
    s2.save()
    ok("salva mantiene la password", fitz.open(P("protetto.pdf")).needs_pass)
    s2.undo()
    s2.save()
    ok("dopo annulla resta protetto", fitz.open(P("protetto.pdf")).needs_pass)
    s2.save_unlocked(P("sbloccato.pdf"))
    d = fitz.open(P("sbloccato.pdf"))
    ok("copia senza password", not d.needs_pass and not d.is_encrypted and "Modulo" in d[0].get_text())
    s2.remove_password()
    s2.save()
    ok("togli password al documento", not fitz.open(P("protetto.pdf")).needs_pass)
    d = fitz.open()
    p = d.new_page()
    p.draw_rect((50, 50, 200, 200), color=(1, 0, 0), fill=(0, 0.6, 1))
    d.save(P("colori.pdf"))
    s3 = Session()
    s3.open(P("colori.pdf"))
    s3.to_grayscale()
    px = s3.doc[0].get_pixmap()
    c = px.pixel(100, 100)
    ok("scala di grigi", c[0] == c[1] == c[2], c)


def t_pdfa_export():
    d = fitz.open()
    p = d.new_page()
    p.insert_text((72, 100), "Testo in Helvetica non incorporato, àèé", fontsize=13)
    p.insert_text((72, 140), "Times", fontsize=13, fontname="tiro")
    p.insert_text((300, 400), "Ruotato", fontsize=13, rotate=90)
    a = p.add_text_annot((400, 100), "nota")
    d.set_metadata({"title": ""})
    d.save(P("non-pdfa.pdf"))
    s = Session()
    s.open(P("non-pdfa.pdf"))
    pre = s.preflight()
    ok("verifica trova i problemi", not pre["pdfa"] and any(i["label"] == "Caratteri incorporati" and not i["ok"] for i in pre["items"]))
    r = s.save_pdfa(P("archivio.pdf"))
    ok("PDF/A senza problemi residui", r["ok"], r["left"])
    a = fitz.open(P("archivio.pdf"))
    txt = a[0].get_text()
    ok("PDF/A: testo intatto", "Helvetica non incorporato, àèé" in txt and "Times" in txt and "Ruotato" in txt, txt)
    ok("PDF/A: niente caratteri mancanti", not any(f[1] == "n/a" for f in a[0].get_fonts()), a[0].get_fonts())
    ok("PDF/A: metadati", "pdfaid:part>2" in a.get_xml_metadata() and a.metadata["title"] == "archivio")
    s.export_text(P("testo.txt"))
    ok("esporta testo", "Times" in open(P("testo.txt"), encoding="utf-8").read())
    s.export_html(P("pagina.html"))
    ok("esporta HTML", "Helvetica" in open(P("pagina.html"), encoding="utf-8").read())
    s2 = Session()
    s2.open(os.path.join(ROOT, "samples", "esempio.pdf"))
    r = s2.export_pptx(P("slide.pptx"))
    from pptx import Presentation
    prs = Presentation(P("slide.pptx"))
    texts = " ".join(sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame)
    ok("esporta PowerPoint", len(prs.slides) == len(s2.doc) and "Relazione" in texts, texts[:80])


def t_signatures():
    c = signing.create_certificate("Mario Rossi", "mario@example.com", "", "prova1234")
    ok("certificato creato", os.path.exists(c["path"]))
    info = signing.certificate_info(c["path"], "prova1234")
    ok("info certificato", info["name"] == "Mario Rossi" and info["selfSigned"])
    try:
        signing.certificate_info(c["path"], "sbagliata")
        ok("password certificato errata", False)
    except RuntimeError as e:
        ok("password certificato errata", "errata" in str(e))
    s = Session()
    s.open(P("form.pdf"))
    s.add_field(0, "signature", [300, 600, 500, 660], "Firma")
    s.save(P("da-firmare.pdf"))
    r = s.sign_digital(P("firmato.pdf"), c["path"], "prova1234", 0, [60, 700, 260, 760], reason="Approvazione", location="Milano")
    ok("firma digitale", os.path.exists(r["path"]) and r["signer"] == "Mario Rossi")
    s2 = Session()
    s2.open(P("firmato.pdf"))
    v = s2.verify_signatures()["signatures"]
    ok("verifica firma", len(v) == 1 and v[0]["intact"] and not v[0]["changedAfter"], v)
    # second signature into the empty field: the first stays valid
    s2.sign_digital(P("firmato2.pdf"), c["path"], "prova1234", field="Firma")
    s3 = Session()
    s3.open(P("firmato2.pdf"))
    v = s3.verify_signatures()["signatures"]
    ok("seconda firma nel campo", len(v) == 2 and all(x["intact"] for x in v), [(x["field"], x["summary"]) for x in v])
    ok("campo firma risulta firmato", any(f["kind"] == "signature" and f["signed"] for f in s3.fields()))
    # tampering
    data = bytearray(open(P("firmato.pdf"), "rb").read())
    k = data.find(b"Modulo di iscrizione")
    if k < 0:                                   # compressed: change a byte inside the first stream
        k = data.find(b"stream") + 20
    data[k] = data[k] ^ 1
    open(P("manomesso.pdf"), "wb").write(data)
    s4 = Session()
    s4.open(P("manomesso.pdf"))
    v = s4.verify_signatures()["signatures"]
    ok("manomissione rilevata", v and not v[0]["intact"], v and v[0]["summary"])


def t_repair_batch():
    data = open(P("form.pdf"), "rb").read()
    cut = data[: data.rfind(b"xref")]               # no cross-reference table
    open(P("rotto.pdf"), "wb").write(cut + b"garbage")
    r = files.repair(P("rotto.pdf"), P("riparato.pdf"))
    ok("ripara", r["repaired"] and r["pages"] == 4 and "Modulo" in fitz.open(P("riparato.pdf"))[0].get_text(), r)
    out = os.path.join(WORK, "lotto")
    os.makedirs(out)
    srcs = [P("form.pdf"), P("colori.pdf"), P("protetto-lotto.pdf")]
    s = Session()
    s.open(P("form.pdf"))
    s.save_protected(srcs[2], "x", "", True, False, False)
    for action, opts in (("compress", {}), ("numbers", {}), ("watermark", {"text": "COPIA"}), ("rotate", {"angle": 90}),
                         ("pdfa", {}), ("grayscale", {}), ("flatten", {}), ("protect", {"newPassword": "pw"}), ("repair", {})):
        r = batch.run(srcs, action, opts, out)
        ok(f"lotto {action}", r["ok"] == 2 and not r["results"][2]["ok"], [(x["name"], x.get("error", x.get("note"))) for x in r["results"]])
    r = batch.run(srcs[2:], "unlock", {"password": "x"}, out)
    ok("lotto rimuovi password", r["ok"] == 1 and not fitz.open(r["results"][0]["out"]).needs_pass)


def t_rotated():
    d = fitz.open()
    d.new_page(width=595, height=842).set_rotation(90)
    d.save(P("ruotata.pdf"))
    s = Session()
    s.open(P("ruotata.pdf"))
    r = s.add_field(0, "text", [100, 50, 300, 80], "T")
    ok("campo su pagina ruotata", [round(v) for v in r["fields"][0]["rect"]] == [100, 50, 300, 80], r["fields"][0]["rect"])
    r = s.add_link(0, [100, 100, 300, 130], uri="a.it")
    ok("collegamento su pagina ruotata", [round(v) for v in r["links"][0]["rect"]] == [100, 100, 300, 130], r["links"][0]["rect"])
    c = signing.create_certificate("Rot", "", "", "abcd1")
    s.sign_digital(P("ruotata-firmata.pdf"), c["path"], "abcd1", 0, [400, 300, 600, 380])
    s2 = Session()
    s2.open(P("ruotata-firmata.pdf"))
    sig = [[round(v) for v in f["rect"]] for f in s2.fields() if f["kind"] == "signature"]
    ok("firma su pagina ruotata dove trascinata", sig == [[400, 300, 600, 380]], sig)


def t_audit():
    """Bugs found by the audit of 1.5.1: they must not come back."""
    s = Session()
    s.open(os.path.join(ROOT, "samples", "esempio.pdf"))
    s.add_shape(0, "rect", [40, 160], [430, 260], "#1e40af")
    s.add_shape(0, "arrow", [480, 300], [420, 230])
    ts = s.tables(0)
    ok("un rettangolo sopra la tabella non la rompe", len(ts) == 1 and len(ts[0]["rows"]) == 4 and len(ts[0]["rows"][0]) == 4,
       [(len(t["rows"]), len(t["rows"][0])) for t in ts])
    d = fitz.open()
    p = d.new_page()
    p.insert_text((72, 72), "Titolo", fontsize=20)
    w = fitz.Widget()
    w.field_type, w.rect, w.field_name, w.field_value = fitz.PDF_WIDGET_TYPE_TEXT, fitz.Rect(100, 200, 300, 220), "a", "ValoreCampo"
    p.add_widget(w)
    d.save(P("campo.pdf"))
    s2 = Session()
    s2.open(P("campo.pdf"))
    ok("il testo di un campo non è un blocco modificabile", [b["text"] for b in s2.blocks(0)["blocks"]] == ["Titolo"])
    ok("il valore di un campo non è un titolo", all(b["title"] != "ValoreCampo" for b in s2.bookmarks_from_headings()))
    # many sensitive data: fast, and the same datum is found once
    d = fitz.open()
    for i in range(60):
        q = d.new_page()
        for k in range(40):
            q.insert_text((40, 60 + k * 18), f"IBAN IT60X0542811101000000123456 tel 3331234567 m{k}@example.com", fontsize=9)
    d.save(P("sensibili.pdf"))
    s3 = Session()
    s3.open(P("sensibili.pdf"))
    import time
    t = time.time()
    hits = s3.find_sensitive(["iban", "email", "phone"])
    dt = time.time() - t
    ok("dati sensibili: 7200 trovati in pochi secondi", len(hits) == 7200 and dt < 10, f"{len(hits)} in {dt:.1f}s")
    # undo of a page keeps the bookmarks; the undo copies stay within their memory
    s3.set_bookmarks([{"level": 1, "title": "Inizio", "page": 0}, {"level": 1, "title": "Pagina 3", "page": 2}])
    s3.add_text(2, 100, 800, "x", {"family": "Arial", "size": 12, "color": "#000000"})
    ok("modifica di una pagina: annulla salva solo la pagina", s3.undo_stack[-1][0] == "page")
    s3.undo()
    ok("annulla di una pagina tiene i segnalibri", [b["title"] for b in s3.bookmarks()] == ["Inizio", "Pagina 3"])
    s3.MAX_UNDO_BYTES = 200 * 1024
    for k in range(6):
        s3.stamp_pages({"footer": f"giro {k}"})
    ok("memoria dell'annulla limitata", sum(Session._snap_size(x) for x in s3.undo_stack) <= 200 * 1024 or len(s3.undo_stack) == 1,
       len(s3.undo_stack))
    sc = fitz.open()
    pg = sc.new_page()
    from PIL import Image as Im
    buf = io.BytesIO()
    Im.new("RGB", (800, 1100), (200, 200, 190)).save(buf, "JPEG")
    pg.insert_image(pg.rect, stream=buf.getvalue())
    sc.save(P("scan1.pdf"))
    s4 = Session()
    s4.open(P("scan1.pdf"))
    ok("formato automatico: scansione in JPEG, testo in PNG", s4.render(0, 1, "auto").startswith("data:image/jpeg")
       and s2.render(0, 1, "auto").startswith("data:image/png"))


make_form()
for name, fn in (("audit", t_audit), ("pagine ruotate", t_rotated), ("moduli", t_forms), ("pagine", t_pages), ("segnalibri/link/proprietà", t_bookmarks_links_props),
                 ("sicurezza/colore", t_security_color), ("pdfa/esporta", t_pdfa_export), ("firme", t_signatures),
                 ("ripara/lotto", t_repair_batch)):
    check(name, fn)
print("TUTTO OK" if not failed else f"{failed} PROBLEMI")
sys.exit(1 if failed else 0)
