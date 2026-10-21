"""Checks that need a real Windows: the built-in OCR on the sample and the
self-update swap script. Usage: python tools/test_windows.py  (exit code 1 on failure)"""
import hashlib
import os
import shutil
import subprocess
import sys
import tempfile
import time

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.environ["CALAMO_DATA"] = tempfile.mkdtemp(prefix="calamo-data-")
sys.path.insert(0, os.path.join(ROOT, "app"))

failures = []


def check(name, ok, detail=""):
    print(("PASS " if ok else "FAIL ") + name + (f"  ({detail})" if detail else ""), flush=True)
    if not ok:
        failures.append(name)


def test_ocr():
    from engine.ocr import available_languages
    from engine.document import Session
    langs = available_languages()
    check("OCR languages available", bool(langs), ", ".join(langs))
    if not langs:
        return
    s = Session()
    s.open(os.path.join(ROOT, "samples", "esempio.pdf"))
    res = s.analyze(0)
    for l in res["ink"]:
        print(f"   [{l['kind']}] {l['source']} doubt={l['doubt']}: {l['text']}")
    hand = [l for l in res["ink"] if l["kind"] == "hand"]
    check("handwriting lines found", len(hand) >= 2, f"{len(hand)} lines")
    check("glyphs learned", len(res["hand"]["glyphs"]) >= 5, f"{len(res['hand']['glyphs'])} glyphs")
    # mark the letters of a line by hand and rewrite it with digits
    if hand:
        L = s.ink_letters(0, hand[0]["id"])
        check("letter boxes proposed", len(L["segments"]) >= 3, f"{len(L['segments'])} boxes")
        segs = [dict(x, char=x["char"] or "x") for x in L["segments"]]
        r = s.teach_letters(0, hand[0]["id"], segs)
        check("letters stored", r["stored"] >= 3, f"{r['stored']} stored")
        s.edit_ink(0, hand[-1]["id"], "Firmato 12/05/2026")
    out = os.path.join(tempfile.mkdtemp(), "out.pdf")
    s.save(out)
    check("save after edits", os.path.getsize(out) > 1000)


def test_rotated_page_edit():
    """Phone scans often have /Rotate: rewriting handwriting must replace it in place."""
    import io
    import numpy as np
    import pymupdf as fitz
    from PIL import Image, ImageDraw, ImageFont
    from engine.document import Session, _view_matrix
    font = ImageFont.truetype(os.path.join(os.environ["WINDIR"], "Fonts", "Inkfree.ttf"), 64)
    img = Image.new("RGB", (900, 110), "white")
    ImageDraw.Draw(img).text((10, 10), "Firmato Giulia", font=font, fill=(25, 40, 110))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    for rot in (0, 90, 270):
        doc = fitz.open()
        p = doc.new_page(width=595, height=842)
        p.set_cropbox(fitz.Rect(20, 30, 580, 820))
        p.set_rotation(rot)
        target = fitz.Rect(60, 200, 330, 233)                     # where it appears on screen
        p.insert_image(target * ~_view_matrix(p), stream=buf.getvalue(), rotate=rot)
        path = os.path.join(tempfile.mkdtemp(), "rot.pdf")
        doc.save(path)
        s = Session()
        s.open(path)
        ink = s.analyze(0)["ink"]
        if not ink:
            check(f"rotation {rot}: handwriting found", False)
            continue
        line = fitz.Rect(ink[0]["bbox"])
        check(f"rotation {rot}: line found where it is shown", target.contains((line.tl + line.br) / 2),
              str([round(v) for v in line]))
        s.edit_ink(0, 0, "Anna")
        page = s.doc[0]

        def ink_share(rect):    # dark pixels in a displayed area, as the user sees the page
            pix = page.get_pixmap(clip=rect, dpi=144)
            a = np.frombuffer(pix.samples, np.uint8).reshape(pix.h, pix.w, pix.n)[:, :, :3].mean(axis=2)
            return float((a < 140).mean())
        start = fitz.Rect(line.x0, line.y0 - 4, line.x0 + line.width * 0.4, line.y1 + 4)
        # where the old writing was and the new, shorter one does not reach
        new_end = s.ink[0][0]["bbox"][2] + 2
        tail = fitz.Rect(max(line.x0 + line.width * 0.6, new_end), line.y0, line.x1, line.y1)
        elsewhere = fitz.Rect(page.rect.x0, line.y1 + line.height * 2, page.rect.x1, page.rect.y1)
        check(f"rotation {rot}: new writing on the old line", ink_share(start) > 0.02, f"{ink_share(start):.3f}")
        left = ink_share(tail) if tail.width > 3 else 0.0      # the new text may cover it all
        check(f"rotation {rot}: old writing covered", left < 0.01, f"{left:.3f}")
        check(f"rotation {rot}: nothing written elsewhere", ink_share(elsewhere) < 0.001, f"{ink_share(elsewhere):.4f}")


def test_writings():
    """A scanned page with print and two people's handwriting: every line must land in
    the writing of its author, and rewriting must keep size, baseline and letters."""
    import io
    import numpy as np
    import pymupdf as fitz
    from PIL import Image, ImageDraw, ImageFilter, ImageFont
    from engine.document import Session
    from engine.handwriting import measure_line
    fonts = os.path.join(os.environ["WINDIR"], "Fonts")
    rows = [("print", "times.ttf", 44, (20, 20, 20), "Il budget approvato resta invariato"),
            ("print", "times.ttf", 44, (20, 20, 20), "Presenti tutti i responsabili di area"),
            ("print", "times.ttf", 44, (20, 20, 20), "Totale previsto per il trimestre"),
            ("A", "Inkfree.ttf", 60, (25, 40, 140), "Consegna entro venerdi"),
            ("A", "Inkfree.ttf", 60, (25, 40, 140), "Firmato Giulia Rossi"),
            ("B", "segoesc.ttf", 54, (170, 25, 25), "Visto e approvato"),
            ("B", "segoesc.ttf", 54, (170, 25, 25), "Marco Bianchi")]
    W, H = 1654, 2339
    scan = Image.new("RGB", (W, H), (247, 246, 241))
    d = ImageDraw.Draw(scan)
    y, truth = 200, []
    for who, f, size, col, text in rows:
        font = ImageFont.truetype(os.path.join(fonts, f), size)
        d.text((180, y), text, font=font, fill=col)
        truth.append((who, text, y * 72 / 200))
        y += int(size * 2.1)
    scan = scan.filter(ImageFilter.GaussianBlur(0.7))
    buf = io.BytesIO()
    scan.save(buf, "JPEG", quality=85)
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_image(page.rect, stream=buf.getvalue())
    path = os.path.join(tempfile.mkdtemp(), "writings.pdf")
    doc.save(path)
    s = Session()
    s.open(path)
    r = s.analyze(0)
    names = {st["id"]: st for st in r["styles"]}
    found = {}
    for l in r["ink"]:
        who = min(truth, key=lambda t: abs(t[2] - l["bbox"][1]))[0]
        found.setdefault(who, []).append(l)
        print(f"   {who:5} -> {names[l['style']]['name']:12} [{l['kind']}] {l['source']}: {l['text']}")
    for who in ("print", "A", "B"):
        check(f"writings: lines of {who} found", bool(found.get(who)))
    if not all(found.get(w) for w in ("print", "A", "B")):
        return
    sid = {w: {l["style"] for l in found[w]} for w in found}
    check("writings: print kept together and printed", len(sid["print"]) == 1
          and all(l["kind"] == "print" for l in found["print"]), str(sid["print"]))
    check("writings: each person is one writing", len(sid["A"]) == 1 and len(sid["B"]) == 1,
          f"A={sid['A']} B={sid['B']}")
    check("writings: the two people are not mixed", not (sid["A"] & sid["B"]) and not (sid["print"] & (sid["A"] | sid["B"])),
          f"A={sid['A']} B={sid['B']}")
    check("writings: handwriting recognised as handwriting",
          all(l["kind"] == "hand" for l in found["A"] + found["B"]))
    # rewrite a printed line: same baseline and size as the original
    pl = found["print"][0]
    base0, xh0 = measure_line(s.doc[0], pl["bbox"], pl["text"])
    s.edit_ink(0, pl["id"], "Il budget approvato aumenta")
    base1, xh1 = measure_line(s.doc[0], s.ink[0][pl["id"]]["bbox"], "Il budget approvato aumenta")
    check("print rewrite: same baseline", abs(base1 - base0) < xh0 * 0.25, f"{base0:.1f} -> {base1:.1f}")
    check("print rewrite: same size", abs(xh1 - xh0) < xh0 * 0.15, f"{xh0:.1f} -> {xh1:.1f}")
    check("print rewrite: still searchable", bool(s.doc[0].search_for("aumenta")))
    # rewrite a line of B: it must be drawn in B's ink colour (its own letters), not A's
    bl = found["B"][0]
    s.edit_ink(0, bl["id"], "Visto, approvato")
    pix = s.doc[0].get_pixmap(clip=fitz.Rect(bl["bbox"]), dpi=144)
    a = np.frombuffer(pix.samples, np.uint8).reshape(pix.h, pix.w, pix.n)[:, :, :3].astype(int)
    ink = a[a.mean(axis=2) < 170]
    red = len(ink) and (ink[:, 0] - ink[:, 2]).mean() > 40
    check("handwriting rewrite: written with its own writer's ink", bool(red), f"{len(ink)} ink px")


def _sha(p):
    with open(p, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def test_portable_swap():
    """The update script must wait for the app to exit, then replace the exe."""
    import updater
    tmp = tempfile.mkdtemp(prefix="calamo upd à-")      # space and accent in the path on purpose
    target = os.path.join(tmp, "CalamoPDF.exe")
    shutil.copy(os.path.join(os.environ["WINDIR"], "System32", "where.exe"), target)
    upd_dir = os.path.join(tmp, "update")
    os.makedirs(upd_dir)
    new = os.path.join(upd_dir, "CalamoPDF.exe")
    shutil.copy(os.path.join(os.environ["WINDIR"], "System32", "whoami.exe"), new)
    want = _sha(new)
    # a fake running app: keeps the target open for 4 s, then exits
    holder = subprocess.Popen([sys.executable, "-c",
                               f"import time; f=open({target!r},'rb'); time.sleep(4)"])
    script = os.path.join(upd_dir, "aggiorna-calamo.cmd")
    with open(script, "w", encoding="utf-8", newline="") as f:
        f.write(updater.build_script({"kind": "portable"}, new, False, target, [holder.pid]))
    t0 = time.time()
    subprocess.Popen(["cmd.exe", "/c", script], creationflags=0x08000000 | 0x00000200)
    time.sleep(1.5)
    check("exe untouched while the app runs", _sha(target) != want)
    while time.time() - t0 < 40 and (not os.path.exists(target) or _sha(target) != want):
        time.sleep(0.5)
    check("exe replaced after the app exits", os.path.exists(target) and _sha(target) == want,
          f"{time.time() - t0:.1f}s")
    time.sleep(1)
    check("update script deletes itself", not os.path.exists(script))


def _cer(a, b):
    """Character error rate of a reading against the truth."""
    import difflib
    a, b = a.lower().strip(), b.lower().strip()
    sm = difflib.SequenceMatcher(None, a, b)
    return 1 - sm.ratio() if b else float(bool(a))


def _line_image(text, font_file, size, colour):
    from PIL import Image, ImageDraw, ImageFont
    font = ImageFont.truetype(os.path.join(os.environ["WINDIR"], "Fonts", font_file), size)
    l, t, r, b = font.getbbox(text)
    img = Image.new("RGB", (r - l + 40, b - t + 30), (247, 246, 241))
    ImageDraw.Draw(img).text((20 - l, 15 - t), text, font=font, fill=colour)
    return img


def test_handreader():
    """The dedicated handwriting reader: download, then read lines Windows OCR finds hard."""
    import io
    import numpy as np
    from engine.handreader import READER
    from engine.ocr import windows_ocr
    check("reader: onnxruntime and tokenizers present", READER.runtime_ok())
    t0 = time.time()
    for attempt in range(3):
        try:
            READER.download()
            break
        except Exception as e:
            print("   download attempt failed:", e)
            time.sleep(5)
    check("reader: model downloaded", READER.installed(), f"{time.time() - t0:.0f} s")
    if not READER.installed():
        return
    lines = [("Consegna entro venerdi", "Inkfree.ttf"), ("Firmato Giulia Rossi", "Inkfree.ttf"),
             ("Ottimo lavoro del team", "Inkfree.ttf"), ("Visto e approvato", "segoesc.ttf"),
             ("Marco Bianchi", "segoesc.ttf"), ("Riunione del 12 maggio", "segoesc.ttf"),
             ("Chiamare il fornitore domani", "Inkfree.ttf"), ("Pagato in contanti", "segoesc.ttf")]
    ocr_cer, model_cer, times = [], [], []
    for text, f in lines:
        img = _line_image(text, f, 56, (25, 40, 140))
        buf = io.BytesIO()
        img.save(buf, "PNG")
        o = " ".join(w["text"] for line in windows_ocr(buf.getvalue()) for w in line)
        t = time.time()
        m, conf = READER.read(np.asarray(img))
        times.append(time.time() - t)
        ocr_cer.append(_cer(o, text))
        model_cer.append(_cer(m, text))
        print(f"   {text!r:32} OCR={o!r:32} lettore={m!r} ({conf:.2f}, {times[-1]:.2f}s)")
    oc, mc = float(np.mean(ocr_cer)), float(np.mean(model_cer))
    check("reader: reads handwriting well", mc < 0.25, f"errore lettore {mc:.2f}, OCR Windows {oc:.2f}")
    check("reader: reasonably fast", float(np.median(times)) < 3.0, f"{np.median(times):.2f} s per riga")


def test_reader_in_analysis():
    """A whole page read with the reader on: the dictionary and the reader must improve
    (never worsen) what Windows OCR reads, and keep the writings apart."""
    import numpy as np
    from engine.document import Session
    from engine.handreader import READER
    if not READER.installed():
        check("reader in analysis: model available", False)
        return
    path, truth = _scan_page(skew=0)
    for reader in (None, READER):
        s = Session()
        s.open(path)
        t = time.time()
        r = s.analyze(0, reader=reader)
        el = time.time() - t
        errs = []
        for l in r["ink"]:
            exp = min(truth, key=lambda x: abs(x[2] - l["bbox"][1]))
            errs.append(_cer(l["text"], exp[1]))
            print(f"   {'lettore' if reader else 'solo OCR'}: {l['source']:7} {l['text']!r} (atteso {exp[1]!r})")
        print(f"   {'lettore' if reader else 'solo OCR'}: errore medio {np.mean(errs):.3f}, {el:.1f} s")
        if reader is None:
            base = float(np.mean(errs))
        else:
            check("reader in analysis: not worse than OCR alone", np.mean(errs) <= base + 0.02,
                  f"{base:.3f} -> {np.mean(errs):.3f}")
            check("reader in analysis: two writers kept apart",
                  len({st["id"] for st in r["styles"] if st["kind"] == "hand"}) >= 2)


def _scan_page(skew=0.0):
    import io
    import pymupdf as fitz
    from PIL import Image, ImageDraw, ImageFilter, ImageFont
    fonts = os.path.join(os.environ["WINDIR"], "Fonts")
    rows = [("print", "times.ttf", 44, (20, 20, 20), "Il budget approvato resta invariato"),
            ("print", "times.ttf", 44, (20, 20, 20), "Presenti tutti i responsabili di area"),
            ("A", "Inkfree.ttf", 60, (25, 40, 140), "Consegna entro venerdi"),
            ("A", "Inkfree.ttf", 60, (25, 40, 140), "Firmato Giulia Rossi"),
            ("B", "segoesc.ttf", 54, (170, 25, 25), "Visto e approvato"),
            ("B", "segoesc.ttf", 54, (170, 25, 25), "Marco Bianchi")]
    W, H = 1654, 2339
    scan = Image.new("RGB", (W, H), (247, 246, 241))
    d = ImageDraw.Draw(scan)
    y, truth = 200, []
    for who, f, size, col, text in rows:
        d.text((180, y), text, font=ImageFont.truetype(os.path.join(fonts, f), size), fill=col)
        truth.append((who, text, y * 72 / 200))
        y += int(size * 2.1)
    if skew:
        scan = scan.rotate(skew, resample=Image.BICUBIC, fillcolor=(247, 246, 241))
    scan = scan.filter(ImageFilter.GaussianBlur(0.7))
    buf = io.BytesIO()
    scan.save(buf, "JPEG", quality=85)
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_image(page.rect, stream=buf.getvalue())
    path = os.path.join(tempfile.mkdtemp(), "scan.pdf")
    doc.save(path)
    return path, truth


def test_skewed_scan():
    """A scan tilted by a few degrees is straightened before reading."""
    import numpy as np
    from engine.document import Session
    for skew in (0, 2.5, -4):
        path, truth = _scan_page(skew)
        s = Session()
        s.open(path)
        r = s.analyze(0)
        errs = [min(_cer(l["text"], t[1]) for t in truth) for l in r["ink"] if l["source"] != "unread"]
        found = len([e for e in errs if e < 0.3])
        print(f"   inclinazione {skew}: {[l['text'] for l in r['ink']]}")
        check(f"skewed scan {skew} deg: lines read", found >= 4, f"{found} righe lette bene su {len(truth)}")


def test_table_form():
    """A handwritten form with drawn borders (slightly crooked, like a photo): every cell
    must be read on its own, never glued to the next cell or to the border."""
    import io
    import numpy as np
    import pymupdf as fitz
    from PIL import Image, ImageDraw, ImageFilter, ImageFont
    from engine.document import Session
    from engine.handreader import READER
    fonts = os.path.join(os.environ["WINDIR"], "Fonts")
    font = ImageFont.truetype(os.path.join(fonts, "Inkfree.ttf"), 46)
    rows = [("1", "Setup ambiente", "Alta", "20/12/2025", "Da fare"),
            ("2", "Database", "Media", "30/12/2025", "In corso"),
            ("3", "Deploy su server", "Bassa", "05/02/2026", "Da fare")]
    W, H = 1654, 1100
    img = Image.new("RGB", (W, H), (244, 243, 238))
    d = ImageDraw.Draw(img)
    xs, y0, rh = [80, 200, 760, 1000, 1330, 1580], 150, 150
    ink = (30, 40, 150)
    for k in range(len(rows) + 1):              # rows of the grid, a little crooked
        d.line([(xs[0], y0 + k * rh), (xs[-1], y0 + k * rh + 9)], fill=ink, width=3)
    for x in xs:
        d.line([(x, y0), (x + 6, y0 + len(rows) * rh)], fill=ink, width=3)
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            d.text((xs[c] + 25, y0 + r * rh + 45), text, font=font, fill=ink)
    img = img.filter(ImageFilter.GaussianBlur(0.6))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=85)
    doc = fitz.open()
    page = doc.new_page(width=595, height=396)
    page.insert_image(page.rect, stream=buf.getvalue())
    path = os.path.join(tempfile.mkdtemp(), "table.pdf")
    doc.save(path)
    s = Session()
    s.open(path)
    r = s.analyze(0, reader=READER if READER.installed() else None)
    texts = [l["text"] for l in r["ink"]]
    print("   tabella:", texts)
    cells = [c for row in rows for c in row if len(c) > 2]
    good = sum(1 for c in cells if any(_cer(t, c) < 0.25 for t in texts))
    check("table: cells read one by one", good >= len(cells) * 0.75, f"{good} di {len(cells)}")
    glued = [t for t in texts if sum(1 for c in set(cells) if c.lower() in t.lower()) >= 2]
    check("table: no cells glued together", not glued, str(glued))


def test_rewrite_keeps_background():
    """Rewriting a line changes only its letters: the highlighter goes on, the table and
    the notebook rules stay, and the new text stays inside its cell."""
    import io
    import numpy as np
    import pymupdf as fitz
    from PIL import Image, ImageDraw, ImageFilter, ImageFont
    from engine.document import Session
    font = ImageFont.truetype(os.path.join(os.environ["WINDIR"], "Fonts", "Inkfree.ttf"), 58)
    W, H = 1654, 760
    img = Image.new("RGB", (W, H), (246, 245, 239))
    d = ImageDraw.Draw(img, "RGBA")
    for y in range(80, H, 105):                                  # notebook rules
        d.line([(0, y), (W, y + 3)], fill=(150, 175, 220), width=3)
    d.rectangle([110, 175, 820, 250], fill=(255, 235, 60, 150))  # highlighter
    xs, y0, rh = [90, 860, 1250, 1580], 150, 180
    ink = (30, 45, 140)
    for i in range(3):
        d.line([(xs[0], y0 + i * rh), (xs[-1], y0 + i * rh + 4)], fill=ink, width=4)
    for x in xs:
        d.line([(x, y0), (x + 3, y0 + 2 * rh)], fill=ink, width=4)
    rows = [("Consegna entro venerdi", "Alta"), ("Firmato Giulia Rossi", "Media")]
    for r, (a, b) in enumerate(rows):
        d.text((xs[0] + 30, y0 + r * rh + 40), a, font=font, fill=ink)
        d.text((xs[1] + 30, y0 + r * rh + 40), b, font=font, fill=ink)
    img = img.filter(ImageFilter.GaussianBlur(0.6))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=88)
    k = 595 / W
    doc = fitz.open()
    page = doc.new_page(width=595, height=H * k)
    page.insert_image(page.rect, stream=buf.getvalue())
    path = os.path.join(tempfile.mkdtemp(), "bg.pdf")
    doc.save(path)
    s = Session()
    s.open(path)
    r = s.analyze(0)
    print("   righe:", [l["text"] for l in r["ink"]])

    def find(words, area):
        l = next((l for l in s._ink_public(0) if any(w in l["text"].lower() for w in words)), None)
        if l is None:      # the OCR missed it among the lines of the table: outline it by hand
            res = s.mark_ink(0, [v * k for v in area])
            l = next((x for x in res["ink"] if x["id"] in res.get("added", [])), None)
            print("   segnata a mano:", l and l["text"])
        return l
    first = find(("venerdi", "consegna", "entro"), (xs[0] + 20, y0 + 25, xs[1] - 10, y0 + rh - 25))
    second = find(("rossi", "firmato", "giulia"), (xs[0] + 20, y0 + rh + 25, xs[1] - 10, y0 + 2 * rh - 25))
    if not first or not second:
        check("background: lines found", False)
        return

    def crop(rect):
        pix = s.doc[0].get_pixmap(dpi=144, clip=fitz.Rect(rect))
        return np.frombuffer(pix.samples, np.uint8).reshape(pix.h, pix.w, pix.n)[:, :, :3].astype(int)
    b = fitz.Rect(first["bbox"])
    s.edit_ink(0, first["id"], "Ok")
    tail = crop([b.x0 + b.width * 0.4, b.y0 + 2, b.x1, b.y1 - 2])
    yellow = ((tail[:, :, 0] > 200) & (tail[:, :, 1] > 190) & (tail[:, :, 2] < 170)).mean()
    check("background: highlighter continues", yellow > 0.85, f"{yellow:.0%}")
    check("background: old letters gone", (tail.mean(axis=2) < 120).mean() < 0.003, f"{(tail.mean(axis=2) < 120).mean():.3%}")
    border = crop([xs[1] * k - 3, b.y0, xs[1] * k + 4, b.y1])
    check("background: table border kept", (border.mean(axis=2) < 150).any(axis=1).mean() > 0.9)
    b2 = fitz.Rect(second["bbox"])
    s.edit_ink(0, second["id"], "Firmato Giulia Rossi e Marco Bianchi per delega del direttore")
    past = crop([xs[1] * k + 2, b2.y0, xs[1] * k + 8, b2.y1])
    check("background: long text stays in its cell", (past.mean(axis=2) < 110).mean() < 0.02)
    under = crop([0, (y0 + 2 * rh) * k + 3, W * k, (y0 + 2 * rh) * k + 40])
    check("background: nothing written below the table", (under.mean(axis=2) < 110).mean() < 0.002)


def test_remove_and_mark():
    """'Not writing' removes a line for good; outlining missed writing reads it."""
    from engine.document import Session
    path, truth = _scan_page(skew=0)
    s = Session()
    s.open(path)
    r = s.analyze(0)
    target = next(l for l in r["ink"] if l["kind"] == "print")
    n = len(r["ink"])
    r = s.remove_ink(0, target["id"])
    check("remove: line removed", len(r["ink"]) == n - 1)
    r = s.analyze(0)
    check("remove: not back after a new reading",
          not any(abs(l["bbox"][1] - target["bbox"][1]) < 3 and abs(l["bbox"][0] - target["bbox"][0]) < 3 for l in r["ink"]))
    # outline the area of the first printed line again: it is read and added
    b = target["bbox"]
    r = s.mark_ink(0, [b[0] - 4, b[1] - 4, b[2] + 4, b[3] + 4])
    new = [l for l in r["ink"] if l["id"] in r.get("added", [])]
    exp = min(truth, key=lambda t: abs(t[2] - b[1]))[1]
    print("   segnata:", [l["text"] for l in new], "atteso:", exp)
    check("mark: outlined writing read", bool(new) and _cer(new[0]["text"], exp) < 0.2,
          new[0]["text"] if new else "nessuna riga")


def test_editing():
    """1.4 editing tools on Windows paths: pages, marks, signature, redaction, stamps, protect."""
    import io
    import pymupdf as fitz
    from PIL import Image, ImageDraw
    from engine.document import Session
    tmp = tempfile.mkdtemp(prefix="calamo-edit-")
    src = os.path.join(tmp, "contratto.pdf")
    doc = fitz.open()
    for i in range(3):
        p = doc.new_page()
        p.insert_text((72, 90), f"Pagina {i + 1}: Mario Rossi RSSMRA80A01H501U mario.rossi@example.com", fontsize=11)
    doc.save(src)
    other = os.path.join(tmp, "allegato.pdf")
    d2 = fitz.open()
    d2.new_page().insert_text((72, 90), "Allegato A", fontsize=14)
    d2.save(other)
    s = Session()
    s.open(src)
    s.merge_pdf(other, after=0)
    s.move_page(1, 3)
    order = [s.doc[i].get_text()[:9] for i in range(4)]
    check("edit: merge and reorder pages", order[3].startswith("Allegato") and len(s.doc) == 4, str(order))
    r = s.mark_text(0, [60, 75, 400, 95], "highlight", "#ffd400")
    check("edit: highlight on text", any(o["type"] == "Highlight" for o in r["objects"]))
    s.add_note(0, 500, 60, "Controllare", "#ffd400")
    check("edit: comments listed", len(s.comments()) == 2)
    img = Image.new("RGB", (400, 160), (238, 235, 228))
    ImageDraw.Draw(img).line([(30, 120), (120, 40), (220, 120), (360, 50)], fill=(20, 20, 90), width=6)
    buf = io.BytesIO()
    img.save(buf, "JPEG")
    s.save_signature("Prova", buf.getvalue(), clean=True)
    r = s.place_signature(0, "Prova", 300, 700)
    check("edit: signature placed", any(o["type"] == "signature" for o in r["objects"]))
    hits = s.find_sensitive(["cf", "email"], "")
    s.redact([{"page": h["page"], "bbox": h["bbox"]} for h in hits], True)
    left = "".join(s.doc[i].get_text() for i in range(len(s.doc)))
    check("edit: sensitive data really removed", "RSSMRA" not in left and "example.com" not in left, f"{len(hits)} hits")
    s.stamp_pages({"footer": "Pagina {n} di {N}", "watermark": "BOZZA"})
    check("edit: page numbers", "Pagina 4 di 4" in s.doc[3].get_text())
    out = os.path.join(tmp, "protetto.pdf")
    s.save_protected(out, "segreta", "", True, False, False)
    d = fitz.open(out)
    check("edit: password protected copy", d.needs_pass and d.authenticate("segreta") > 0)
    light = os.path.join(tmp, "leggero.pdf")
    r = s.save_optimized(light, "media")
    check("edit: lighter copy opens", len(fitz.open(light)) == 4, f"{r['before']} -> {r['after']} byte")
    cmp = s.compare_with(src)
    check("edit: compare finds changes", cmp["changed"] >= 1, f"{cmp['changed']} di {cmp['total']}")
    # drag a block of text
    b0 = next(b for b in s.blocks(0)["blocks"] if "Mario" in b["text"])
    r = s.move_block(0, b0["bbox"], 30, 200)
    b1 = next((b for b in r["blocks"] if "Mario" in b["text"] and abs(b["bbox"][0] - b0["bbox"][0] - 30) < 2 and b["bbox"][1] > b0["bbox"][1] + 150), None)
    check("drag: text moved with its font", b1 is not None, b1 and str([round(v) for v in b1["bbox"]]))
    # text written after a save stays readable (saving used to leave stale font references)
    s.save(os.path.join(tmp, "salvato.pdf"))
    s.add_text(0, 72, 400, "Scritto dopo il salvataggio", {"family": "Arial", "size": 12})
    b2 = next((b for b in s.blocks(0)["blocks"] if b["bbox"][1] > 380 and b["bbox"][1] < 420), None)
    check("save: text written afterwards is readable", b2 is not None and b2["text"] == "Scritto dopo il salvataggio",
          b2 and b2["text"])
    check("save: spaces are real spaces (copy and search)", bool(s.doc[0].search_for("dopo il salvataggio")))
    # merge and split whole files
    from engine import files as FILES
    merged = os.path.join(tmp, "unito.pdf")
    FILES.merge([{"src": src, "name": "contratto.pdf"}, {"src": other, "name": "allegato.pdf", "pages": "1"}], merged)
    check("files: merged with one bookmark per file", len(fitz.open(merged)) == 4 and len(fitz.open(merged).get_toc()) == 2)
    folder = os.path.join(tmp, "parti")
    os.makedirs(folder)
    r = FILES.split(merged, "unito", folder, "ranges", "1-2, 3-4")
    check("files: split by ranges", [len(fitz.open(f)) for f in r["files"]] == [2, 2], str([os.path.basename(f) for f in r["files"]]))


def test_more():
    """1.4: OCR layer with the real Windows OCR, conversions without Office, tables and
    translation with Windows fonts."""
    import io
    import pymupdf as fitz
    from PIL import Image
    from engine.document import Session
    from engine import convert as C
    tmp = tempfile.mkdtemp(prefix="calamo-more-")
    # searchable: read the sample with Windows OCR, then the writing can be searched
    s = Session()
    s.open(os.path.join(ROOT, "samples", "esempio.pdf"))
    res = s.analyze(0)
    words = [l["text"].split()[0] for l in res["ink"] if l["kind"] == "hand" and l["text"].split()]
    r = s.make_searchable(0)
    hit = s.doc[0].search_for(words[0]) if words else []
    check("ocr: handwriting searchable after OCR", r["written"] > 0 and bool(hit), f"{r['written']} parole, cerco {words[:1]}")
    # conversions: images and Word without Office (built-in or LibreOffice)
    img = os.path.join(tmp, "foto.jpg")
    Image.new("RGB", (900, 600), (180, 40, 40)).save(img)
    out = os.path.join(tmp, "foto.pdf")
    C.files_to_pdf([img], out)
    check("convert: image to PDF", len(fitz.open(out)) == 1)
    import docx
    d = docx.Document()
    d.add_heading("Relazione dopo il collaudo", 1)
    d.add_paragraph("Affitto e ufficio.")
    dp = os.path.join(tmp, "rel.docx")
    d.save(dp)
    out = os.path.join(tmp, "rel.pdf")
    how = C.to_pdf(dp, out)
    page = fitz.open(out)[0]
    check("convert: Word to PDF", bool(page.search_for("dopo il collaudo")) or "collaudo" in page.get_text(), how)
    imgs = C.pdf_to_images(fitz.open(out), [0], tmp, "rel", "png", 72)
    check("convert: PDF to images", os.path.exists(imgs["files"][0]))
    # a table rewritten with the Windows fonts
    doc = fitz.open()
    p = doc.new_page()
    X, Y = [72, 220, 360], [100, 125, 150, 175]
    for x in X:
        p.draw_line((x, Y[0]), (x, Y[-1]))
    for y in Y:
        p.draw_line((X[0], y), (X[-1], y))
    for r_, row in enumerate([["Voce", "Importo"], ["Affitto", "1200"], ["Luce", "80"]]):
        for c_, t in enumerate(row):
            p.insert_text((X[c_] + 4, Y[r_] + 17), t, fontsize=10)
    tp = os.path.join(tmp, "tab.pdf")
    doc.save(tp)
    s = Session()
    s.open(tp)
    s.edit_table(0, 0, [{"r": 1, "c": 0, "text": "Affitto ufficio"}])
    s.table_row(0, 0, 2, "add")
    cells = s.table_cells(0, 0)["cells"]
    check("table: cell edited and row added", len(cells) == 4 and cells[1][0] == "Affitto ufficio", str(cells))
    check("table: spaces real in the new text", bool(s.doc[0].search_for("Affitto ufficio")))
    # translation keeps the layout (a fake translator: the service is online)
    s.translate_into("en", lambda texts, tgt: [t.replace("Voce", "Item").replace("Luce", "Electricity") for t in texts])
    check("translate: text rewritten in place", any("Electricity" in b["text"] for b in s.blocks(0)["blocks"]))


def test_graphics():
    """1.4.2: every element of a page can be dragged: a table with its text, a drawing, a
    picture, a form field, and on a scan a graphic of the image (here a stamp)."""
    import io
    import pymupdf as fitz
    from PIL import Image, ImageDraw
    from engine.document import Session
    tmp = tempfile.mkdtemp(prefix="calamo-gfx-")
    d = fitz.open()
    p = d.new_page(width=595, height=842)
    for i in range(4):
        p.draw_line((60, 100 + i * 20), (300, 100 + i * 20), width=1)
    for x in (60, 180, 300):
        p.draw_line((x, 100), (x, 160), width=1)
    for i in range(3):
        p.insert_text((65, 115 + i * 20), f"Voce {i + 1}", fontsize=10)
    p.draw_circle((450, 120), 25, color=(0.8, 0, 0), fill=(1, 0.8, 0))
    buf = io.BytesIO()
    Image.new("RGB", (80, 50), (30, 120, 200)).save(buf, "PNG")
    p.insert_image(fitz.Rect(60, 300, 140, 350), stream=buf.getvalue())
    w = fitz.Widget()
    w.field_type, w.field_name, w.rect, w.field_value = fitz.PDF_WIDGET_TYPE_TEXT, "nome", fitz.Rect(300, 300, 450, 320), "Mario"
    p.add_widget(w)
    path = os.path.join(tmp, "g.pdf")
    d.save(path)
    for rot in (0, 90):
        s = Session()
        s.open(path)
        if rot:
            s.rotate_page(0, rot)
        g = s.graphics(0)
        kinds = sorted(e["kind"] for e in g)
        check(f"graphics {rot}: drawings, picture and field found", kinds == ["field", "image", "vector", "vector"], kinds)
        for e in g:
            s.move_graphic(0, e, 20, 150)
        after = s.graphics(0)
        moved = all(any(a["kind"] == e["kind"] and abs(a["bbox"][1] - min(e["bbox"][1] + 150, s.doc[0].rect.height - (e["bbox"][3] - e["bbox"][1]))) < 2
                        for a in after) for e in g)
        check(f"graphics {rot}: every element moved", moved, [(e["kind"], [round(v) for v in e["bbox"]]) for e in after])
        words = [w_[4] for w_ in s.doc[0].get_text("words")]
        check(f"graphics {rot}: the table took its text along", words.count("Voce") == 3, words)
        check(f"graphics {rot}: nothing left behind", len(s.doc[0].get_drawings()) == len(fitz.open(path)[0].get_drawings()))
    # a scan with a stamp
    img = Image.new("RGB", (1240, 1754), (238, 234, 225))
    dr = ImageDraw.Draw(img)
    dr.ellipse((850, 1350, 1100, 1600), outline=(190, 30, 40), width=8)
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=90)
    d = fitz.open()
    d.new_page(width=595, height=842).insert_image(fitz.Rect(0, 0, 595, 842), stream=buf.getvalue())
    path = os.path.join(tmp, "scan.pdf")
    d.save(path)
    s = Session()
    s.open(path)
    g = [e for e in s.graphics(0) if e["kind"] == "scan"]
    check("scan graphic found", len(g) == 1, g)
    if g:
        s.move_graphic(0, g[0], -250, -400)
        pix = s.doc[0].get_pixmap(dpi=72)
        a = pix.samples
        def red(x, y):
            o = (y * pix.width + x) * pix.n
            return a[o] > 150 and a[o + 1] < 110
        old = sum(red(x, y) for x in range(400, 535) for y in range(640, 775))
        new = sum(red(x, y) for x in range(150, 285) for y in range(240, 375))
        check("scan graphic lifted and laid down elsewhere", old < 5 and new > 50, f"prima {old}, dopo {new}")


def test_real_services():
    """What only a real Windows (and the internet) has: online translation, Office
    printing to PDF, the voices for reading aloud. Skipped, not failed, where missing."""
    from engine import extras as EX
    try:
        out = EX.translate_texts(["Buongiorno a tutti", "Il contratto è firmato."], "en")
        check("translate (online): Italian to English", "morning" in out[0].lower() and "contract" in out[1].lower(), out)
    except RuntimeError as e:
        print("SKIP translate (online):", e)
    if sys.platform != "win32":
        return
    import winreg
    try:
        winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "Word.Application").Close()
        word = True
    except OSError:
        word = False
    if word:
        import docx
        import pymupdf as fitz
        from engine import convert as C
        tmp = tempfile.mkdtemp(prefix="calamo-office-")
        d = docx.Document()
        d.add_paragraph("Stampato da Microsoft Word")
        dp = os.path.join(tmp, "w.docx")
        d.save(dp)
        out = os.path.join(tmp, "w.pdf")
        check("Office: Word prints to PDF", C._office_com(dp, out) and "Microsoft Word" in fitz.open(out)[0].get_text())
    else:
        print("SKIP Office: Microsoft Word is not installed here")
    r = subprocess.run(["powershell", "-NoProfile", "-Command",
                        "Add-Type -AssemblyName System.Speech; (New-Object System.Speech.Synthesis.SpeechSynthesizer)"
                        ".GetInstalledVoices() | ForEach-Object { $_.VoiceInfo.Name + ' ' + $_.VoiceInfo.Culture }"],
                       capture_output=True, text=True, timeout=60)
    voices = [v for v in r.stdout.splitlines() if v.strip()]
    print("   voices:", voices)
    if voices:
        r = subprocess.run(["powershell", "-NoProfile", "-Command",
                            "Add-Type -AssemblyName System.Speech; $s = New-Object System.Speech.Synthesis.SpeechSynthesizer; "
                            "$f = Join-Path $env:TEMP 'calamo-voce.wav'; $s.SetOutputToWaveFile($f); "
                            "$s.Speak('Buongiorno, questa è una prova'); $s.Dispose(); (Get-Item $f).Length"],
                           capture_output=True, text=True, timeout=60)
        size = int((r.stdout.strip().splitlines() or ["0"])[-1] or 0)
        check("voices: a system voice speaks", size > 10000, f"{size} byte")
    else:
        print("SKIP voices: no voice installed on this machine")


def _scan_form(path, rows=5):
    """A photographed form: a table drawn slightly slanting, with printed text in its cells."""
    import io
    import pymupdf as fitz
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (1240, 1754), (238, 235, 226))
    d = ImageDraw.Draw(img)
    font = None
    for f in ("arial.ttf", "C:/Windows/Fonts/arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
        try:
            font = ImageFont.truetype(f, 34)
            break
        except Exception:
            continue
    d.text((120, 120), "Elenco attività", fill=(25, 25, 25), font=font)
    xs = [100, 420, 900, 1140]
    ys = [260 + k * 110 for k in range(rows + 1)]
    for y in ys:
        d.line((xs[0], y, xs[-1], y + 8), fill=(20, 20, 20), width=4)
    for x in xs:
        d.line((x, ys[0], x + 5, ys[-1]), fill=(20, 20, 20), width=4)
    data = [("Voce", "Descrizione", "Stato")] + [(f"Voce {k}", f"Attivita numero {k}", "Fatto") for k in range(1, rows)]
    for r, row in enumerate(data):
        for c, text in enumerate(row):
            d.text((xs[c] + 25, ys[r] + 35), text, fill=(25, 25, 25), font=font)
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=90)
    doc = fitz.open()
    doc.new_page(width=595, height=842).insert_image(fitz.Rect(0, 0, 595, 842), stream=buf.getvalue())
    doc.save(path)
    return data


def test_export_scan():
    """1.5.2: a scanned page exports as text and tables, not as one picture of the page."""
    import docx
    import numpy as np
    import pymupdf as fitz
    from engine.document import Session
    from engine import export
    from engine.layout import scan_tables
    tmp = tempfile.mkdtemp(prefix="calamo-exp-")
    path = os.path.join(tmp, "modulo.pdf")
    data = _scan_form(path)
    s = Session()
    s.open(path)
    pix = s.doc[0].get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
    rgb = np.frombuffer(pix.samples, np.uint8).reshape(pix.height, pix.width, 3)
    t = scan_tables(rgb)
    check("scan table found (rows, columns)", len(t) == 1 and len(t[0]["h"]) - 1 == len(data) and len(t[0]["v"]) - 1 == 3,
          [(len(x["h"]) - 1, len(x["v"]) - 1) for x in t])
    if sys.platform != "win32":
        return                                       # the reading needs the Windows OCR
    s.analyze(0)
    out = os.path.join(tmp, "modulo.docx")
    export.to_docx(s, out)
    d = docx.Document(out)
    pics = len(d.inline_shapes)
    check("export scan: no picture of the whole page", pics == 0, f"{pics} immagini")
    cells = [[c.text for c in r.cells] for r in d.tables[0].rows] if d.tables else []
    check("export scan: the table as a Word table", len(cells) == len(data) and len(cells[0]) == 3, cells[:2])
    ok = sum(1 for r in range(min(len(cells), len(data))) for c in range(3)
             if c < len(cells[r]) and data[r][c].lower().replace("à", "a") in cells[r][c].lower().replace("à", "a"))
    check("export scan: cells read in their place", ok >= len(data) * 3 * 0.7, f"{ok} di {len(data) * 3}: {cells}")
    check("export scan: title as text", any("Elenco" in p.text for p in d.paragraphs), [p.text for p in d.paragraphs][:3])


def test_print_now():
    """1.5.2: quick print without any dialog (to a file, with Microsoft Print to PDF)."""
    if sys.platform != "win32":
        return
    import pymupdf as fitz
    from engine import printing
    tmp = tempfile.mkdtemp(prefix="calamo-print-")
    doc = fitz.open()
    doc.new_page().insert_text((72, 100), "Pagina uno da stampare", fontsize=20)
    p2 = doc.new_page(width=842, height=595)
    p2.insert_text((72, 100), "Pagina due, orizzontale", fontsize=20)
    out = os.path.join(tmp, "stampa.pdf")
    print("   stampante predefinita:", printing.default_printer())
    try:
        printing.print_pages(doc, [0, 1], "Prova", printer="Microsoft Print to PDF", output=out)
    except RuntimeError as e:
        print("SKIP print:", e)
        return
    for _ in range(40):                            # the spooler writes the file a moment later
        if os.path.exists(out) and os.path.getsize(out) > 1000:
            break
        time.sleep(0.5)
    n = len(fitz.open(out)) if os.path.exists(out) else 0
    check("print now: two pages printed", n == 2, f"{n} pagine")


def test_setup_script():
    import updater
    text = updater.build_script({"kind": "setup"}, r"C:\T\CalamoPDF-Setup.exe", True,
                                r"C:\Program Files\Calamo PDF Studio\CalamoPDF.exe", [1, 2])
    check("setup script is silent and relaunches", "/VERYSILENT" in text and "/relaunch=1" in text)
    check("setup script keeps the install folder", '/DIR="C:\\Program Files\\Calamo PDF Studio"' in text)


def test_release_check():
    import updater
    try:
        info = updater.check("0.0.1")
        check("GitHub release found", bool(info and info.get("url")), info and info.get("name"))
    except Exception as e:
        check("GitHub release found", False, str(e))


if __name__ == "__main__":
    for t in (test_ocr, test_writings, test_skewed_scan, test_handreader, test_reader_in_analysis, test_table_form, test_rewrite_keeps_background, test_remove_and_mark,
              test_rotated_page_edit, test_editing, test_more, test_graphics, test_real_services, test_export_scan, test_print_now, test_portable_swap,
              test_setup_script, test_release_check):
        try:
            t()
        except Exception as e:
            import traceback
            traceback.print_exc()
            check(t.__name__, False, repr(e))
    print("\n%d failure(s)" % len(failures))
    sys.exit(1 if failures else 0)
