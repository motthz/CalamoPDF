"""Export to Word (.docx) and Excel (.xlsx), keeping styles, images and tables,
plus the recognised handwriting."""
import io
import re

import pymupdf as fitz
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import fonts as F

HAND_FONT = "Segoe Print"


def _tables(page):
    try:
        if page.first_annot is not None:
            # without the comments: a rectangle drawn over a table is not one of its lines
            one = fitz.open()
            one.insert_pdf(page.parent, from_page=page.number, to_page=page.number, annots=False, links=False)
            _tables.keep = one                  # the tables read their text from this page
            page = one[0]
        return [t for t in page.find_tables().tables if t.row_count >= 2 and t.col_count >= 2]
    except Exception:
        return []


def _inside(rect, rects):
    c = fitz.Point((rect.x0 + rect.x1) / 2, (rect.y0 + rect.y1) / 2)
    return any(r.contains(c) for r in rects)


def _ink_lines(ink, pi):
    """Recognised lines to export (handwriting and scanned print, as read or rewritten)."""
    return list(ink.get(pi, []))


def _hidden_rects(page):
    """Invisible text: a scanner's OCR layer or the searchable copy of rewritten print.
    On a page Calamo has read, the recognised lines replace it in the export."""
    out = []
    for b in page.get_text("dict")["blocks"]:
        for l in b.get("lines", []):
            for s in l["spans"]:
                if s["text"].strip() and s.get("alpha", 255) == 0:
                    out.append(fitz.Rect(s["bbox"]))
    return out


# --------------------------------------------------------------------------- Word
def _font_size(line):
    printed = line.get("kind") == "print"
    return max(8, min(round((line["bbox"][3] - line["bbox"][1]) * (0.75 if printed else 0.8)), 28))


def to_docx(session, path, progress=None):
    """A Word document with the text as paragraphs (with their font, size and colour),
    tables as Word tables, pictures as pictures. On a scanned page the lines Calamo read
    become text, its ruled tables Word tables, its drawings pictures: never the page as
    one picture (unless there is nothing to read on it, like a photo)."""
    from .layout import page_items
    doc = session.doc
    word = Document()
    st = word.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(0)
    first = doc[0].rect
    sec = word.sections[0]
    sec.page_width, sec.page_height = Pt(first.width), Pt(first.height)
    if first.width > first.height:
        sec.orientation = WD_ORIENT.LANDSCAPE
    margin = Pt(40)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = margin
    usable = first.width - 80

    n = len(doc)
    for pi in range(n):
        if progress:
            progress(pi, n)
        items, info = page_items(session, pi)
        W = info["size"][0]
        last_bottom = None
        row_para, row_bottom = None, None
        for it in items:
            kind, bb = it["kind"], it["bbox"]
            gap = Pt(max(0, min(bb[1] - last_bottom, 40))) if last_bottom is not None else None
            if kind == "text":
                b = it["block"]
                para = word.add_paragraph()
                pf = para.paragraph_format
                if gap is not None:
                    pf.space_before = gap
                indent = bb[0] - 40
                if indent > 6:
                    pf.left_indent = Pt(min(indent, usable * 0.6))
                lines = b["lines"]
                if len(lines) == 1:
                    cx = (bb[0] + bb[2]) / 2
                    if abs(cx - W / 2) < 12 and bb[2] - bb[0] < W * 0.7:
                        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        pf.left_indent = None
                for li, l in enumerate(lines):
                    for sp in l["spans"]:
                        if not sp["text"]:
                            continue
                        run = para.add_run(sp["text"])
                        bold, italic, serif, mono = F.style_from(F.clean_name(sp["font"]), sp["flags"])
                        run.bold, run.italic = bold, italic
                        run.font.size = Pt(round(sp["size"] * 2) / 2)
                        run.font.name = F.css_family(sp["font"], sp["flags"])
                        c = sp["color"]
                        run.font.color.rgb = RGBColor((c >> 16) & 255, (c >> 8) & 255, c & 255)
                        if sp["flags"] & 1:
                            run.font.superscript = True
                    if li < len(lines) - 1:
                        last_text = "".join(sp["text"] for sp in l["spans"]).rstrip()
                        if last_text.endswith("-") and para.runs:
                            para.runs[-1].text = para.runs[-1].text.rstrip()[:-1]
                        else:
                            para.add_run(" ")
                last_bottom, row_para = bb[3], None
            elif kind == "ink":
                l = it["line"]
                printed = l.get("kind") == "print"
                # lines side by side on one row of the page stay on one row, separated by a tab
                if row_para is not None and bb[0] >= row_bottom[2] - 2 and (
                        min(bb[3], row_bottom[3]) - max(bb[1], row_bottom[1])) >= 0.5 * min(bb[3] - bb[1], row_bottom[3] - row_bottom[1]):
                    para = row_para
                    para.add_run("\t")
                else:
                    para = word.add_paragraph()
                    if gap is not None:
                        para.paragraph_format.space_before = gap
                    indent = bb[0] - 40
                    if indent > 6:
                        para.paragraph_format.left_indent = Pt(min(indent, usable * 0.6))
                run = para.add_run(l["text"])
                run.font.name = "Calibri" if printed else HAND_FONT
                run.font.size = Pt(_font_size(l))
                run.font.color.rgb = RGBColor(0x1A, 0x1B, 0x20) if printed else RGBColor(0x1F, 0x2A, 0x5A)
                row_para, row_bottom = para, list(bb)
                last_bottom = max(last_bottom or 0, bb[3])
            elif kind == "table":
                rows = it["rows"]
                cols = max(len(r) for r in rows)
                tbl = word.add_table(rows=len(rows), cols=cols)
                tbl.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci, val in enumerate(row):
                        cell = tbl.cell(ri, ci)
                        cell.text = (val or "").replace("\n", " ")
                        if ri == 0:
                            for p in cell.paragraphs:
                                for run in p.runs:
                                    run.bold = True
                last_bottom, row_para = bb[3], None
            elif kind == "image":
                try:
                    w = min(bb[2] - bb[0], usable)
                    word.add_picture(io.BytesIO(it["png"]), width=Pt(w))
                    word.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                except Exception:
                    pass
                last_bottom, row_para = bb[3], None
        if pi < n - 1:
            word.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    word.save(path)
    return path


# --------------------------------------------------------------------------- Excel
_num = re.compile(r"^-?[\d.]+(,\d+)?$|^-?[\d,]+(\.\d+)?$")


def _cell_value(v):
    if v is None:
        return None
    s = str(v).replace("\n", " ").strip()
    t = s.replace("€", "").replace("%", "").replace(" ", "")
    if t and _num.match(t):
        try:
            if "," not in t and re.match(r"^-?\d{1,3}(\.\d{3})+$", t):
                t = t.replace(".", "")          # Italian thousands: 120.400
            if "," in t and "." in t:
                t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
            elif "," in t:
                t = t.replace(",", ".")
            val = float(t)
            if s.endswith("%"):
                return round(val / 100, 10)
            return int(val) if val.is_integer() and "." not in t else val
        except ValueError:
            pass
    return s


def _layout_rows(page, ink_lines):
    """No ruled tables: rebuild a grid from word positions."""
    words = page.get_text("words")
    if ink_lines:
        hidden = _hidden_rects(page)
        words = [w for w in words if not _inside(fitz.Rect(w[:4]), hidden)]
    items = [(w[0], w[1], w[2], w[3], w[4]) for w in words]
    for l in ink_lines:
        items.append((l["bbox"][0], l["bbox"][1], l["bbox"][2], l["bbox"][3], l["text"]))
    if not items:
        return []
    items.sort(key=lambda w: ((w[1] + w[3]) / 2, w[0]))
    rows, cur, cy = [], [], None
    for w in items:
        mid = (w[1] + w[3]) / 2
        if cy is None or abs(mid - cy) <= max(3, (w[3] - w[1]) * 0.45):
            cur.append(w)
            cy = mid if cy is None else (cy + mid) / 2
        else:
            rows.append(cur)
            cur, cy = [w], mid
    if cur:
        rows.append(cur)
    # column anchors: cluster cell starting x positions
    cells_per_row = []
    for row in rows:
        row.sort(key=lambda w: w[0])
        cells, cell = [], [row[0]]
        for w in row[1:]:
            h = max(w[3] - w[1], 4)
            if w[0] - cell[-1][2] > h * 1.2:
                cells.append(cell)
                cell = [w]
            else:
                cell.append(w)
        cells.append(cell)
        cells_per_row.append([(c[0][0], " ".join(x[4] for x in c)) for c in cells])
    anchors = sorted({round(x / 12) * 12 for row in cells_per_row for x, _ in row})
    merged = []
    for a in anchors:
        if not merged or a - merged[-1] > 24:
            merged.append(a)
    out = []
    for row in cells_per_row:
        line = [None] * len(merged)
        for x, text in row:
            ci = min(range(len(merged)), key=lambda k: abs(merged[k] - x))
            line[ci] = (line[ci] + " " + text) if line[ci] else text
        out.append(line)
    return out


def to_xlsx(session, path, progress=None):
    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="1F2430")
    head_font = Font(bold=True, color="F4D58D")
    thin = Side(style="thin", color="C9CED8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    zebra = PatternFill("solid", fgColor="F5F6F9")
    count = 0

    def write(ws, rows, header=True):
        widths = {}
        for ri, row in enumerate(rows, start=1):
            for ci, val in enumerate(row, start=1):
                v = _cell_value(val)
                c = ws.cell(row=ri, column=ci, value=v)
                c.border = border
                c.alignment = Alignment(vertical="top", wrap_text=True)
                if header and ri == 1:
                    c.fill, c.font = head_fill, head_font
                elif ri % 2 == 0:
                    c.fill = zebra
                if isinstance(v, float) and val and str(val).strip().endswith("%"):
                    c.number_format = "0.00%"
                widths[ci] = max(widths.get(ci, 8), min(len(str(v or "")) + 2, 60))
        for ci, w in widths.items():
            ws.column_dimensions[get_column_letter(ci)].width = w
        if header and rows:
            ws.freeze_panes = "A2"

    from .layout import page_items
    doc, ink = session.doc, session.ink
    for pi in range(len(doc)):
        page = doc[pi]
        if progress:
            progress(pi, len(doc))
        items, _ = page_items(session, pi)
        tables = [it["rows"] for it in items if it["kind"] == "table"]
        for ti, rows in enumerate(tables, start=1):
            ws = wb.create_sheet(f"P{pi + 1} Tabella {ti}"[:31])
            write(ws, rows, header=True)
            count += 1
        if not tables:
            rows = _layout_rows(page, _ink_lines(ink, pi))
            if rows:
                ws = wb.create_sheet(f"Pagina {pi + 1}")
                write(ws, rows, header=False)
                count += 1
    if not wb.sheetnames:
        wb.create_sheet("Vuoto")
    wb.save(path)
    return count
