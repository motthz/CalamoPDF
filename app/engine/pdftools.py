"""The PDF tools of the big suites: forms (fill, create, recognise, flatten), cropping and
resizing pages, several pages per sheet and booklets, bookmarks, links, document properties,
removing a password, grayscale, PDF/A for archiving with a check of the document, and
exporting to PowerPoint, HTML and plain text.

Like the rest of the app, coordinates coming from the interface are displayed coordinates
(what the user sees); they are converted to the page's own where PDF objects are made."""
import datetime
import math
import os
import re
from collections import Counter

import numpy as np
import pymupdf as fitz

from . import fonts as F
from .editing import parse_pages

FIELD_TYPES = {"text": fitz.PDF_WIDGET_TYPE_TEXT, "checkbox": fitz.PDF_WIDGET_TYPE_CHECKBOX,
               "radio": fitz.PDF_WIDGET_TYPE_RADIOBUTTON, "combo": fitz.PDF_WIDGET_TYPE_COMBOBOX,
               "list": fitz.PDF_WIDGET_TYPE_LISTBOX, "signature": fitz.PDF_WIDGET_TYPE_SIGNATURE}
FIELD_KIND = {v: k for k, v in FIELD_TYPES.items()}
FIELD_LABEL = {"text": "Testo", "checkbox": "Casella", "radio": "Scelta singola", "combo": "Menu a tendina",
               "list": "Elenco", "signature": "Firma", "other": "Campo"}

# paper sizes in points (portrait)
SIZES = {"A3": (841.89, 1190.55), "A4": (595.28, 841.89), "A5": (419.53, 595.28),
         "Letter": (612.0, 792.0), "Legal": (612.0, 1008.0)}
MM = 72 / 25.4
CHECK_CHARS = "☐□❏◻⬜▢"
PDFA_RGB = "sRGB IEC61966-2.1"


def _size(name, w=None, h=None):
    if name in SIZES:
        return SIZES[name]
    if w and h:
        return float(w) * MM, float(h) * MM
    raise RuntimeError("Formato della pagina non valido.")


def _fit(src, box):
    """src (w, h) scaled to fit box, centred: the rectangle it takes."""
    k = min(box.width / src[0], box.height / src[1])
    w, h = src[0] * k, src[1] * k
    x, y = box.x0 + (box.width - w) / 2, box.y0 + (box.height - h) / 2
    return fitz.Rect(x, y, x + w, y + h)


def _baked(doc):
    """A copy of the document with annotations and form fields drawn into the pages: what
    show_pdf_page needs to place a page with everything that is on it."""
    copy = fitz.open(stream=doc.tobytes(encryption=fitz.PDF_ENCRYPT_NONE), filetype="pdf")
    try:
        copy.bake()
    except Exception:
        pass
    return copy


def _pdf_date(dt):
    return dt.strftime("D:%Y%m%d%H%M%S+00'00'")


def _xmp_escape(s):
    return (s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _overlaps(a, b, k=0.4):
    i = a & b
    if i.is_empty:
        return False
    return abs(i) > k * min(abs(a), abs(b))


class PdfToolsMixin:
    """Mixed into Session: uses self.doc, self.lock, self._checkpoint, self._touch, self.info,
    self.ink, self.hands, self._shown, self._text_rect."""

    def _forget_pages(self, pages):
        """The coordinates of these pages changed: the handwriting read there is read again."""
        for i in pages:
            self.ink.pop(i, None)
            self.hands.pop(i, None)
            self.ignored.pop(i, None)

    # ================================================================== forms
    def _widget(self, page, xref):
        for w in page.widgets() or []:
            if w.xref == int(xref):
                return w
        raise RuntimeError("Campo non trovato (forse è stato eliminato).")

    def _field_public(self, i, page, w):
        kind = FIELD_KIND.get(w.field_type, "other")
        value = w.field_value
        if kind in ("checkbox", "radio"):
            on = w.on_state()
            value = value is True or (value not in (False, None, "", "Off") and (value == on or value == "Yes" or on is True))
        signed = False
        if kind == "signature":
            signed = bool(getattr(w, "is_signed", False))
        return {"page": i, "id": w.xref, "name": w.field_name or "", "label": w.field_label or w.field_name or "",
                "kind": kind, "kindLabel": FIELD_LABEL.get(kind, "Campo"), "value": value,
                "options": list(w.choice_values or []), "rect": list(self._shown(page, w.rect)),
                "readonly": bool((w.field_flags or 0) & 1), "required": bool((w.field_flags or 0) & 2),
                "multiline": kind == "text" and bool((w.field_flags or 0) & 4096), "signed": signed}

    def fields(self, index=None):
        """The form fields of a page, or of the whole document."""
        with self.lock:
            pages = range(len(self.doc)) if index is None else [int(index)]
            out = []
            for i in pages:
                page = self.doc[i]
                for w in page.widgets() or []:
                    out.append(self._field_public(i, page, w))
            return out

    def fill_fields(self, values):
        """values: [{page, id, value}]. Text, checkbox (true / false), choice (the option)."""
        with self.lock:
            if not values:
                return {**self.info(), "fields": self.fields()}
            self._checkpoint()
            touched = set()
            for v in values:
                page = self.doc[int(v["page"])]
                w = self._widget(page, v["id"])
                kind = FIELD_KIND.get(w.field_type, "other")
                val = v.get("value")
                if kind == "checkbox":
                    w.field_value = bool(val)
                elif kind == "radio":
                    if val:
                        # the other buttons of the group go off
                        for i2, p2 in enumerate(self.doc):
                            for o in p2.widgets(types=[fitz.PDF_WIDGET_TYPE_RADIOBUTTON]) or []:
                                if o.field_name == w.field_name and o.xref != w.xref:
                                    o.field_value = False
                                    o.update()
                                    touched.add(i2)
                        w = self._widget(page, v["id"])
                        w.field_value = True
                    else:
                        w.field_value = False
                elif kind == "list" and isinstance(val, list):
                    w.field_value = val
                elif kind == "signature":
                    continue
                else:
                    w.field_value = "" if val is None else str(val)
                w.update()
                touched.add(int(v["page"]))
            self._touch(*touched)
            return {**self.info(), "fields": self.fields()}

    def _field_names(self):
        return {w.field_name for p in self.doc for w in (p.widgets() or []) if w.field_name}

    def _unique_name(self, base, taken):
        base = re.sub(r"[\s.:]+", " ", (base or "").strip()).strip()[:40] or "Campo"
        name, k = base, 2
        while name in taken:
            name = f"{base} {k}"
            k += 1
        taken.add(name)
        return name

    def _make_widget(self, page, kind, rect_unrot, name, options=None, multiline=False, value=None):
        w = fitz.Widget()
        w.field_type = FIELD_TYPES[kind]
        w.rect = fitz.Rect(rect_unrot)
        w.field_name = name
        w.field_label = name
        w.border_color = (0.55, 0.6, 0.75)
        w.border_width = 0.8
        w.fill_color = (0.93, 0.95, 1.0)
        w.text_font = "Helv"
        w.text_fontsize = 0                      # automatic: the text fits the field
        if kind in ("combo", "list"):
            opts = [o for o in (options or []) if str(o).strip()] or ["Opzione 1", "Opzione 2"]
            w.choice_values = opts
            w.field_value = value if value in opts else opts[0]
        elif kind in ("checkbox", "radio"):
            w.field_value = bool(value)
        elif kind == "text":
            w.field_value = value or ""
            if multiline:
                w.field_flags |= 4096
        elif kind == "signature":
            w.fill_color = None
        page.add_widget(w)
        return w

    def add_field(self, index, kind, rect, name="", options=None, multiline=False):
        """A new form field drawn on the page (rect in displayed coordinates)."""
        with self.lock:
            if kind not in FIELD_TYPES:
                raise RuntimeError("Tipo di campo sconosciuto.")
            page = self.doc[int(index)]
            r = fitz.Rect(rect)
            if kind in ("checkbox", "radio") and (r.width < 6 or r.height < 6):
                c = fitz.Point((r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2)
                r = fitz.Rect(c.x - 6, c.y - 6, c.x + 6, c.y + 6)
            if r.width < 6 or r.height < 6:
                raise RuntimeError("Trascina un riquadro un po' più grande per il campo.")
            self._checkpoint()
            taken = self._field_names()
            if kind == "radio" and name and name in taken:
                pass                              # one more button of an existing group
            else:
                name = self._unique_name(name or FIELD_LABEL[kind], taken)
            w = self._make_widget(page, kind, self._text_rect(page, r), name, options, multiline)
            self._touch(int(index))
            return {**self.info(), "fields": self.fields(), "id": w.xref}

    def edit_field(self, index, xref, name=None, options=None, multiline=None, required=None):
        with self.lock:
            page = self.doc[int(index)]
            w = self._widget(page, xref)
            self._checkpoint()
            if name and name.strip() and name.strip() != w.field_name:
                w.field_name = self._unique_name(name, self._field_names())
                w.field_label = w.field_name
            if options is not None and FIELD_KIND.get(w.field_type) in ("combo", "list"):
                opts = [o for o in options if str(o).strip()]
                if opts:
                    w.choice_values = opts
                    if w.field_value not in opts:
                        w.field_value = opts[0]
            if multiline is not None and FIELD_KIND.get(w.field_type) == "text":
                w.field_flags = (w.field_flags | 4096) if multiline else (w.field_flags & ~4096)
            if required is not None:
                w.field_flags = (w.field_flags | 2) if required else (w.field_flags & ~2)
            w.update()
            self._touch(int(index))
            return {**self.info(), "fields": self.fields()}

    def delete_field(self, index, xref):
        with self.lock:
            page = self.doc[int(index)]
            w = self._widget(page, xref)
            self._checkpoint()
            page.delete_widget(w)
            self._touch(int(index))
            return {**self.info(), "fields": self.fields()}

    def detect_fields(self, spec=""):
        """Recognise the places to fill in a form made for paper: dotted and underscored
        blanks, lines to write on, empty boxes and checkboxes. They become real fields,
        named after the label written beside them."""
        with self.lock:
            pages = parse_pages(spec, len(self.doc))
            found = {i: self._detect_on(self.doc[i]) for i in pages}
            total = sum(len(v) for v in found.values())
            if not total:
                return {**self.info(), "fields": self.fields(), "added": 0}
            self._checkpoint()
            taken = self._field_names()
            for i, items in found.items():
                page = self.doc[i]
                for kind, rect, label in items:
                    self._make_widget(page, kind, rect, self._unique_name(label or FIELD_LABEL[kind], taken))
                if items:
                    self._touch(i)
            return {**self.info(), "fields": self.fields(), "added": total}

    def _detect_on(self, page):
        """[(kind, rect in page coordinates, label)] for one page."""
        cands = []
        taken = [fitz.Rect(w.rect) for w in (page.widgets() or [])]
        words = page.get_text("words")
        raw = page.get_text("rawdict", flags=fitz.TEXT_PRESERVE_WHITESPACE)
        text_boxes = []

        def label_left(r):
            """The words just before r on its line: the name of the field."""
            row = [w for w in words if w[2] <= r.x0 + 2 and min(w[3], r.y1) - max(w[1], r.y0) > 0.4 * (w[3] - w[1])
                   and r.x0 - w[2] < 160 and "__" not in w[4] and ".." not in w[4]]
            row.sort(key=lambda w: w[0])
            # the words close together just before the blank
            out = []
            x = r.x0
            for w in reversed(row):
                if x - w[2] > (45 if not out else 12):
                    break
                out.insert(0, w[4])
                x = w[0]
            return " ".join(out).strip(" :;,-.")[:40]

        def label_right(r):
            """The words just after a checkbox: what it means."""
            row = sorted([w for w in words if w[0] >= r.x1 - 1 and min(w[3], r.y1 + 3) - max(w[1], r.y0 - 3) > 0
                          and w[0] - r.x1 < 200], key=lambda w: w[0])
            out, x = [], r.x1
            for w in row:
                if w[0] - x > (20 if not out else 8) or len(out) >= 5:
                    break
                out.append(w[4])
                x = w[2]
            return " ".join(out).strip(" :;,.")[:40]

        def label_above(r):
            above = [w for w in words if w[3] <= r.y0 + 2 and r.y0 - w[3] < 18 and w[0] < r.x1 and w[2] > r.x0]
            above.sort(key=lambda w: (-w[3], w[0]))
            if not above:
                return ""
            y = above[0][3]
            return " ".join(w[4] for w in sorted([w for w in above if abs(w[3] - y) < 3], key=lambda w: w[0])).strip(" :")[:40]

        for b in raw.get("blocks", []):
            if b.get("type") != 0:
                continue
            for line in b["lines"]:
                horizontal = abs(line["dir"][0] - 1) < 1e-3
                for s in line["spans"]:
                    chars = s["chars"]
                    if not chars:
                        continue
                    text_boxes.append(fitz.Rect(s["bbox"]))
                    if not horizontal:
                        continue
                    size = s["size"]
                    base = s["origin"][1]
                    k = 0
                    while k < len(chars):
                        c = chars[k]["c"]
                        if c in CHECK_CHARS:
                            r = fitz.Rect(chars[k]["bbox"])
                            side = min(r.width, r.height, size) * 0.85
                            cx, cy = (r.x0 + r.x1) / 2, base - size * 0.35
                            box = fitz.Rect(cx - side / 2, cy - side / 2, cx + side / 2, cy + side / 2)
                            cands.append(("checkbox", box, label_right(box)))
                            k += 1
                            continue
                        if c in "_.…":
                            j = k
                            while j < len(chars) and chars[j]["c"] == c:
                                j += 1
                            n = j - k
                            if (c == "_" and n >= 4) or (c == "." and n >= 6) or (c == "…" and n >= 2):
                                r = fitz.Rect(chars[k]["bbox"])
                                for ch in chars[k:j]:
                                    r |= fitz.Rect(ch["bbox"])
                                if r.width >= 24:
                                    box = fitz.Rect(r.x0, base - size * 1.05, r.x1, base + size * 0.3)
                                    cands.append(("text", box, label_left(box)))
                            k = j
                            continue
                        k += 1

        # lines to write on and empty boxes, drawn as vectors
        hlines, vlines, boxes = [], [], []
        try:
            drawings = page.get_drawings()
        except Exception:
            drawings = []
        for d in drawings:
            for it in d["items"]:
                if it[0] == "l":
                    p, q = it[1], it[2]
                    if abs(p.y - q.y) < 0.8 and abs(p.x - q.x) >= 50:
                        hlines.append((min(p.x, q.x), max(p.x, q.x), (p.y + q.y) / 2))
                    elif abs(p.x - q.x) < 0.8 and abs(p.y - q.y) > 4:
                        vlines.append(((p.x + q.x) / 2, min(p.y, q.y), max(p.y, q.y)))
                elif it[0] == "re":
                    r = fitz.Rect(it[1])
                    if r.height < 1.6 and r.width >= 50:
                        hlines.append((r.x0, r.x1, (r.y0 + r.y1) / 2))
                    elif r.width < 1.6 and r.height > 4:
                        vlines.append(((r.x0 + r.x1) / 2, r.y0, r.y1))
                    else:
                        filled_dark = d.get("fill") and sum(d["fill"][:3]) < 2.4
                        if not filled_dark:
                            boxes.append(r)
        for r in boxes:
            if 6 <= r.width <= 20 and 6 <= r.height <= 20 and abs(r.width - r.height) < 3:
                cands.append(("checkbox", fitz.Rect(r), label_right(r)))
            elif r.width >= 60 and 12 <= r.height <= 60:
                inner = fitz.Rect(r.x0 + 2, r.y0 + 2, r.x1 - 2, r.y1 - 2)
                if not any(_overlaps(t, inner, 0.2) for t in text_boxes):
                    cands.append(("text", inner, label_left(r) or label_above(r)))
        for x0, x1, y in hlines:
            # a line of a table (with vertical lines at its ends) is not a line to write on
            if any(abs(vx - x0) < 2.5 or abs(vx - x1) < 2.5 for vx, vy0, vy1 in vlines if vy0 - 2 <= y <= vy1 + 2):
                continue
            box = fitz.Rect(x0, y - 15, x1, y - 0.6)
            # a line under some text is an underline: the place is taken
            if any(_overlaps(t, box, 0.25) for t in text_boxes):
                continue
            cands.append(("text", box, label_left(box) or label_above(box)))

        out = []
        for kind, r, label in cands:
            r = r & page.cropbox if page.cropbox.contains(r.tl) else r
            if r.is_empty or r.width < 5 or r.height < 5:
                continue
            if any(_overlaps(r, t, 0.3) for t in taken):
                continue
            taken.append(r)
            out.append((kind, r, label))
            if len(out) >= 120:
                break
        return out

    def flatten(self, forms=True, annots=True):
        """Form fields and / or comments drawn into the pages: they can no longer be
        changed, and look the same everywhere."""
        with self.lock:
            if not forms and not annots:
                raise RuntimeError("Scegli cosa appiattire.")
            self._checkpoint()
            self.doc.bake(annots=bool(annots), widgets=bool(forms))
            self._touch(*range(len(self.doc)))
            return self.info()

    # ================================================================== crop and resize pages
    def _content_box(self, page):
        """Where something is printed on the page (displayed coordinates)."""
        z = 0.75
        pix = page.get_pixmap(matrix=fitz.Matrix(z, z), colorspace=fitz.csGRAY, alpha=False, annots=True)
        a = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width)
        ys, xs = np.nonzero(a < 235)
        if not len(xs):
            return None
        return fitz.Rect(xs.min() / z, ys.min() / z, (xs.max() + 1) / z, (ys.max() + 1) / z)

    def crop_pages(self, spec="", rect=None, margins=None, auto=False):
        """Cut the edges of pages: a rectangle drawn on the page, margins in millimetres
        [top, right, bottom, left], or automatically around what is printed."""
        with self.lock:
            pages = parse_pages(spec, len(self.doc))
            plans = []
            for i in pages:
                page = self.doc[i]
                R = page.rect
                if auto:
                    c = self._content_box(page)
                    if c is None:
                        continue
                    pad = 4 * MM
                    r = fitz.Rect(c.x0 - pad, c.y0 - pad, c.x1 + pad, c.y1 + pad) & R
                elif margins:
                    t, rr, b, l = [max(0.0, float(x or 0)) * MM for x in (list(margins) + [0, 0, 0, 0])[:4]]
                    r = fitz.Rect(l, t, R.width - rr, R.height - b)
                elif rect:
                    r = fitz.Rect(rect) & R
                else:
                    raise RuntimeError("Indica cosa ritagliare.")
                if r.is_empty or r.width < 36 or r.height < 36:
                    raise RuntimeError(f"Il ritaglio della pagina {i + 1} è troppo piccolo.")
                if abs(r.x0) + abs(r.y0) + abs(r.x1 - R.width) + abs(r.y1 - R.height) < 0.5:
                    continue
                plans.append((i, r))
            if not plans:
                return {**self.info(), "cropped": 0}
            self._checkpoint()
            for i, r in plans:
                page = self.doc[i]
                cb = page.cropbox
                u = (r * page.derotation_matrix).normalize()
                u = fitz.Rect(u.x0 + cb.x0, u.y0 + cb.y0, u.x1 + cb.x0, u.y1 + cb.y0) & page.mediabox
                page.set_cropbox(u)
            self._forget_pages([i for i, _ in plans])
            self._touch(*[i for i, _ in plans])
            return {**self.info(), "cropped": len(plans)}

    def reset_crop(self, spec=""):
        with self.lock:
            pages = [i for i in parse_pages(spec, len(self.doc)) if self.doc[i].cropbox != self.doc[i].mediabox]
            if not pages:
                return {**self.info(), "cropped": 0}
            self._checkpoint()
            for i in pages:
                self.doc[i].set_cropbox(self.doc[i].mediabox)
            self._forget_pages(pages)
            self._touch(*pages)
            return {**self.info(), "cropped": len(pages)}

    def _rebuild(self, new):
        """The document is replaced by `new` (same number of pages or not)."""
        try:
            toc = self.doc.get_toc(simple=False)
            if toc:
                new.set_toc([t for t in toc if 1 <= t[2] <= len(new)])
        except Exception:
            pass
        try:
            new.set_metadata({k: v for k, v in self.doc.metadata.items()
                              if k in ("title", "author", "subject", "keywords", "creator", "producer",
                                       "creationDate", "modDate") and v})
        except Exception:
            pass
        data = new.tobytes(garbage=3, deflate=True)
        self.doc = fitz.open(stream=data, filetype="pdf")
        self.resolver = F.FontResolver(self.doc)

    def resize_pages(self, spec="", size="A4", width=None, height=None):
        """Put pages on another paper size, scaled to fit (comments and fields are drawn
        into the page)."""
        with self.lock:
            W, H = _size(size, width, height)
            pages = set(parse_pages(spec, len(self.doc)))
            self._checkpoint()
            src = _baked(self.doc)
            new = fitz.open()
            for i in range(len(self.doc)):
                if i in pages:
                    r = self.doc[i].rect
                    w, h = (H, W) if r.width > r.height else (W, H)
                    p = new.new_page(width=w, height=h)
                    p.show_pdf_page(_fit((r.width, r.height), p.rect), src, i)
                else:
                    new.insert_pdf(self.doc, from_page=i, to_page=i, annots=True, links=True)
            self._rebuild(new)
            self._forget_pages(pages)
            self._touch(*range(len(self.doc)))
            return {**self.info(), "resized": len(pages)}

    def nup_bytes(self, cols=2, rows=1, size="A4", spec="", border=False, landscape=None):
        """Several pages on each sheet: a new document."""
        with self.lock:
            cols, rows = max(1, int(cols)), max(1, int(rows))
            per = cols * rows
            if per < 2:
                raise RuntimeError("Scegli almeno due pagine per foglio.")
            pages = parse_pages(spec, len(self.doc))
            W, H = _size(size)
            if landscape is None:
                first = self.doc[pages[0]].rect
                landscape = (cols > rows) == (first.height >= first.width)
            if landscape:
                W, H = H, W
            src = _baked(self.doc)
            out = fitz.open()
            m, g = 18, 10
            cw, ch = (W - 2 * m - (cols - 1) * g) / cols, (H - 2 * m - (rows - 1) * g) / rows
            sheet = None
            for k, pno in enumerate(pages):
                j = k % per
                if j == 0:
                    sheet = out.new_page(width=W, height=H)
                cx, cy = j % cols, j // cols
                cell = fitz.Rect(m + cx * (cw + g), m + cy * (ch + g), m + cx * (cw + g) + cw, m + cy * (ch + g) + ch)
                r = src[pno].rect
                place = _fit((r.width, r.height), cell)
                sheet.show_pdf_page(place, src, pno)
                if border:
                    sheet.draw_rect(place, color=(0.6, 0.6, 0.6), width=0.5)
            return out.tobytes(garbage=3, deflate=True)

    def booklet_bytes(self, size="A4"):
        """An A5 booklet on A4 sheets (or the half of the chosen size): print on both sides,
        fold in the middle and staple."""
        with self.lock:
            W, H = _size(size)
            W, H = max(W, H), min(W, H)                 # landscape sheet
            src = _baked(self.doc)
            n = len(src)
            N = n + (-n) % 4
            out = fitz.open()
            halves = (fitz.Rect(0, 0, W / 2, H), fitz.Rect(W / 2, 0, W, H))
            for s in range(N // 4):
                for left, right in ((N - 1 - 2 * s, 2 * s), (2 * s + 1, N - 2 - 2 * s)):
                    sheet = out.new_page(width=W, height=H)
                    for pno, half in ((left, halves[0]), (right, halves[1])):
                        if pno < n:
                            r = src[pno].rect
                            box = fitz.Rect(half.x0 + 12, half.y0 + 12, half.x1 - 12, half.y1 - 12)
                            sheet.show_pdf_page(_fit((r.width, r.height), box), src, pno)
            return out.tobytes(garbage=3, deflate=True)

    # ================================================================== bookmarks
    def bookmarks(self):
        with self.lock:
            return [{"level": t[0], "title": t[1], "page": max(0, t[2] - 1)} for t in self.doc.get_toc(simple=True)]

    def set_bookmarks(self, items):
        """items: [{level, title, page}] in order; levels are fixed so the tree is valid."""
        with self.lock:
            n = len(self.doc)
            toc, prev = [], 0
            for it in items or []:
                title = str(it.get("title", "")).strip() or "Senza titolo"
                lvl = max(1, min(int(it.get("level") or 1), prev + 1))
                page = max(0, min(int(it.get("page") or 0), n - 1))
                toc.append([lvl, title, page + 1])
                prev = lvl
            self._checkpoint()
            self.doc.set_toc(toc)
            self._touch()
            return {**self.info(), "bookmarks": self.bookmarks()}

    def bookmarks_from_headings(self):
        """Titles found in the text (bigger or bold short lines), as bookmarks to propose."""
        with self.lock:
            sizes = Counter()
            lines = []
            for i, page in enumerate(self.doc):
                d = page.get_text("dict")
                for b in d["blocks"]:
                    if b.get("type") != 0:
                        continue
                    blines = b["lines"]
                    for l in blines:
                        spans = [s for s in l["spans"] if s["text"].strip()]
                        if not spans:
                            continue
                        text = "".join(s["text"] for s in l["spans"]).strip()
                        for s in spans:
                            sizes[round(s["size"], 1)] += len(s["text"])
                        size = max(s["size"] for s in spans)
                        bold = all(s["flags"] & 16 or "bold" in s["font"].lower() for s in spans)
                        lines.append((i, l["bbox"][1], text, round(size, 1), bold, len(blines)))
            if not sizes:
                return []
            body = sizes.most_common(1)[0][0]
            repeated = Counter(t for _, _, t, _, _, _ in lines)
            npages = len(self.doc)
            cands = []
            for i, y, text, size, bold, nb in lines:
                if len(text) < 3 or len(text) > 90 or re.fullmatch(r"[\d\s.,/-]+", text):
                    continue
                if npages > 2 and repeated[text] > max(2, npages * 0.3):
                    continue                          # running headers and footers
                if size >= body * 1.18 or (bold and size >= body * 0.98 and nb <= 2 and len(text) < 70):
                    cands.append((i, y, text, size, bold))
            if not cands:
                return []
            levels = sorted({c[3] for c in cands}, reverse=True)[:3]
            out = []
            for i, y, text, size, bold in cands:
                lvl = next((k + 1 for k, s in enumerate(levels) if size >= s - 0.05), 3)
                if out and out[-1]["page"] == i and out[-1]["title"] == text:
                    continue
                out.append({"level": lvl, "title": re.sub(r"\s+", " ", text), "page": i})
            low = min(o["level"] for o in out)
            prev = 0
            for o in out:
                o["level"] = max(1, min(o["level"] - low + 1, prev + 1))
                prev = o["level"]
            return out[:300]

    # ================================================================== links
    def links(self, index):
        with self.lock:
            page = self.doc[int(index)]
            out = []
            for l in page.get_links():
                kind = {fitz.LINK_URI: "uri", fitz.LINK_GOTO: "page", fitz.LINK_NAMED: "page"}.get(l["kind"], "other")
                # get_links gives the rectangle as displayed already (insert_link takes it unrotated)
                out.append({"id": l.get("xref", 0), "rect": list(l["from"]), "kind": kind,
                            "uri": l.get("uri", ""), "page": l.get("page", -1) if kind == "page" else -1})
            return out

    def add_link(self, index, rect, uri=None, target_page=None):
        with self.lock:
            page = self.doc[int(index)]
            r = fitz.Rect(rect)
            if r.width < 3 or r.height < 3:
                raise RuntimeError("Trascina un riquadro sopra il testo da collegare.")
            d = {"from": self._text_rect(page, r).normalize()}
            if uri:
                uri = uri.strip()
                if "@" in uri and not re.match(r"^[a-z]+:", uri, re.I):
                    uri = "mailto:" + uri
                elif not re.match(r"^[a-z][a-z0-9+.-]*:", uri, re.I):
                    uri = "https://" + uri
                d.update(kind=fitz.LINK_URI, uri=uri)
            elif target_page is not None:
                tp = int(target_page)
                if not (0 <= tp < len(self.doc)):
                    raise RuntimeError("Questa pagina non esiste.")
                d.update(kind=fitz.LINK_GOTO, page=tp, to=fitz.Point(0, 0))
            else:
                raise RuntimeError("Indica un indirizzo web o una pagina.")
            self._checkpoint()
            page.insert_link(d)
            self.doc.reload_page(page)             # the page keeps its list of links until reloaded
            self._touch(int(index))
            return {**self.info(), "links": self.links(index)}

    def delete_link(self, index, xref):
        with self.lock:
            page = self.doc[int(index)]
            link = next((l for l in page.get_links() if l.get("xref") == int(xref)), None)
            if not link:
                raise RuntimeError("Collegamento non trovato.")
            self._checkpoint()
            page.delete_link(link)
            self.doc.reload_page(page)
            self._touch(int(index))
            return {**self.info(), "links": self.links(index)}

    # ================================================================== properties
    def _lang(self):
        try:
            v = self.doc.xref_get_key(self.doc.pdf_catalog(), "Lang")
            return v[1].strip("()") if v[0] == "string" else ""
        except Exception:
            return ""

    def properties(self):
        with self.lock:
            m = self.doc.metadata or {}
            cat = self.doc.pdf_catalog()
            tagged = self.doc.xref_get_key(cat, "StructTreeRoot")[0] != "null"
            nsig = sum(1 for p in self.doc for w in (p.widgets(types=[fitz.PDF_WIDGET_TYPE_SIGNATURE]) or [])
                       if getattr(w, "is_signed", False))

            def date(s):
                mt = re.match(r"D:(\d{4})(\d{2})?(\d{2})?(\d{2})?(\d{2})?", s or "")
                if not mt:
                    return ""
                y, mo, d, h, mi = mt.groups()
                return f"{d or '01'}/{mo or '01'}/{y}" + (f" {h}:{mi or '00'}" if h else "")
            size = os.path.getsize(self.path) if self.path and os.path.exists(self.path) else None
            return {"title": m.get("title", ""), "author": m.get("author", ""), "subject": m.get("subject", ""),
                    "keywords": m.get("keywords", ""), "creator": m.get("creator", ""), "producer": m.get("producer", ""),
                    "created": date(m.get("creationDate")), "modified": date(m.get("modDate")),
                    "format": m.get("format", ""), "lang": self._lang(), "pages": len(self.doc),
                    "encrypted": bool(getattr(self, "_password", None)) or bool(self.doc.is_encrypted),
                    "forms": sum(1 for p in self.doc for _ in (p.widgets() or [])), "signatures": nsig,
                    "tagged": tagged, "bookmarks": len(self.doc.get_toc()), "size": size,
                    "pageSize": self._paper_name(self.doc[0].rect)}

    @staticmethod
    def _paper_name(r):
        w, h = sorted((r.width, r.height))
        for name, (a, b) in SIZES.items():
            if abs(w - a) < 3 and abs(h - b) < 3:
                return f"{name} {'orizzontale' if r.width > r.height else 'verticale'}"
        return f"{r.width / MM:.0f} × {r.height / MM:.0f} mm"

    def set_properties(self, props):
        with self.lock:
            self._checkpoint()
            m = {k: v for k, v in (self.doc.metadata or {}).items()
                 if k in ("title", "author", "subject", "keywords", "creator", "producer", "creationDate", "modDate", "trapped")}
            for k in ("title", "author", "subject", "keywords"):
                if k in props:
                    m[k] = str(props[k] or "").strip()
            m["modDate"] = _pdf_date(datetime.datetime.now(datetime.timezone.utc))
            self.doc.set_metadata(m)
            if "lang" in props:
                lang = str(props["lang"] or "").strip()
                self.doc.xref_set_key(self.doc.pdf_catalog(), "Lang", fitz.get_pdf_str(lang) if lang else "null")
            self._touch()
            return {**self.info(), "properties": self.properties()}

    # ================================================================== security, colour
    def save_unlocked(self, path):
        """A copy without the password and without the limits on printing and copying."""
        with self.lock:
            data = self.doc.tobytes(encryption=fitz.PDF_ENCRYPT_NONE)
            fitz.open(stream=data, filetype="pdf").save(path, garbage=3, deflate=True, encryption=fitz.PDF_ENCRYPT_NONE)
            return {"path": path}

    def to_grayscale(self):
        """Every colour turned to shades of gray (text, drawings and pictures): for printing
        in black and white."""
        with self.lock:
            self._checkpoint()
            self.doc.recolor(1)
            self._touch(*range(len(self.doc)))
            return self.info()

    # ================================================================== check and PDF/A
    def _missing_fonts(self, doc=None):
        doc = doc or self.doc
        miss = {}
        for i, page in enumerate(doc):
            for f in page.get_fonts():
                if f[1] == "n/a" and f[2] != "Type3":
                    miss.setdefault(F.clean_name(f[3]) or f[4], []).append(i)
        return miss

    def _has_js(self, doc):
        try:
            for x in range(1, min(doc.xref_length(), 30000)):
                try:
                    obj = doc.xref_object(x, compressed=True)
                except Exception:
                    continue
                if "/JS" in obj or "/JavaScript" in obj:
                    return True
        except Exception:
            pass
        return False

    def preflight(self, doc=None):
        """What keeps the document from being a PDF/A for archiving, and from being
        accessible to people who use a screen reader."""
        with self.lock:
            doc = doc or self.doc
            cat = doc.pdf_catalog()
            items = []

            def add(group, ok, label, detail="", fix=""):
                items.append({"group": group, "ok": bool(ok), "label": label, "detail": detail, "fix": fix})
            miss = self._missing_fonts(doc)
            add("pdfa", not miss, "Caratteri incorporati",
                "" if not miss else "Non incorporati: " + ", ".join(sorted(miss)[:6]), "pdfa")
            enc = doc.is_encrypted or (doc is self.doc and bool(getattr(self, "_password", None)))
            add("pdfa", not enc, "Senza password né cifratura", "" if not enc else "Il documento è protetto.", "pdfa")
            js = self._has_js(doc)
            add("pdfa", not js, "Senza JavaScript", "" if not js else "Contiene codice JavaScript.", "pdfa")
            emb = doc.embfile_count()
            add("pdfa", not emb, "Senza file allegati", "" if not emb else f"{emb} file allegati.", "pdfa")
            oi = doc.xref_get_key(cat, "OutputIntents")[0] != "null"
            add("pdfa", oi, "Profilo colore (Output Intent)", "" if oi else "Manca il profilo colore sRGB.", "pdfa")
            xmp = doc.get_xml_metadata() or ""
            ida = "pdfaid:part" in xmp
            add("pdfa", ida, "Dichiarazione PDF/A nei metadati", "" if ida else "Il file non si dichiara PDF/A.", "pdfa")
            hidden = 0
            for p in doc:
                for a in p.annots() or []:
                    if not (a.flags & fitz.PDF_ANNOT_IS_PRINT) or a.flags & (fitz.PDF_ANNOT_IS_HIDDEN | fitz.PDF_ANNOT_IS_NO_VIEW):
                        hidden += 1
            add("pdfa", not hidden, "Commenti stampabili", "" if not hidden else f"{hidden} commenti non stampabili o nascosti.", "pdfa")

            m = doc.metadata or {}
            add("access", bool((m.get("title") or "").strip()), "Titolo del documento",
                "" if m.get("title") else "Il titolo viene mostrato al posto del nome del file.", "props")
            lang = doc.xref_get_key(cat, "Lang")[0] == "string"
            add("access", lang, "Lingua del documento", "" if lang else "Serve ai lettori di schermo per pronunciare il testo.", "props")
            scans = [i + 1 for i, p in enumerate(doc) if not p.get_text("text").strip() and p.get_images()]
            add("access", not scans, "Testo leggibile (non solo immagini)",
                "" if not scans else "Pagine scansionate senza testo: " + ", ".join(map(str, scans[:12])) + ("…" if len(scans) > 12 else ""), "ocr")
            toc = doc.get_toc()
            add("access", bool(toc) or len(doc) < 6, "Segnalibri per navigare",
                "" if toc or len(doc) < 6 else "Un documento lungo è più facile da usare con i segnalibri.", "bookmarks")
            tagged = doc.xref_get_key(cat, "StructTreeRoot")[0] != "null"
            add("access", tagged, "Struttura (tag) per i lettori di schermo",
                "" if tagged else "Il PDF non ha la struttura dei contenuti: va creata nel programma d'origine (Word, ecc.) salvando come PDF accessibile.", "")
            pdfa = all(it["ok"] for it in items if it["group"] == "pdfa")
            return {"items": items, "pdfa": pdfa, "access": all(it["ok"] for it in items if it["group"] == "access")}

    def _embed_missing_fonts(self, doc):
        """Rewrite the text set in fonts that are not in the file with installed fonts that
        get embedded, at the same place, size and width. Returns the pages changed."""
        from .document import REDACT_KW, _append_words
        changed = []
        resolver = F.FontResolver(doc)
        for i, page in enumerate(doc):
            bad = {F.clean_name(f[3]) for f in page.get_fonts() if f[1] == "n/a" and f[2] != "Type3"}
            if not bad:
                continue
            rot = page.rotation
            if rot:
                page.set_rotation(0)
            d = page.get_text("dict", flags=fitz.TEXT_PRESERVE_WHITESPACE | fitz.TEXT_MEDIABOX_CLIP)
            todo = []
            for b in d["blocks"]:
                if b.get("type") != 0:
                    continue
                for l in b["lines"]:
                    if any(F.clean_name(s["font"]) in bad for s in l["spans"]):
                        todo.append(l)
            if not todo:
                if rot:
                    page.set_rotation(rot)
                continue
            for l in todo:
                r = fitz.Rect(l["bbox"])
                # narrowed across the text, so the lines next to it are not touched
                if abs(l["dir"][0]) >= abs(l["dir"][1]):
                    r = fitz.Rect(r.x0, r.y0 + r.height * 0.18, r.x1, r.y1 - r.height * 0.18)
                else:
                    r = fitz.Rect(r.x0 + r.width * 0.18, r.y0, r.x1 - r.width * 0.18, r.y1)
                page.add_redact_annot(r)
            page.apply_redactions(**REDACT_KW)
            for l in todo:
                dx, dy = l["dir"]
                angle = -math.degrees(math.atan2(dy, dx))
                for s in l["spans"]:
                    text = s["text"]
                    if not text.strip():
                        continue
                    name = F.clean_name(s["font"])
                    b_, it_, serif, mono = F.style_from(name, s["flags"])
                    font = None
                    if name not in bad:
                        font = resolver.embedded(page, s["font"])
                        if font is not None and not all(font.has_glyph(ord(c)) for c in text if not c.isspace()):
                            font = None
                    if font is None:
                        font = F.system_font(name, b_, it_, serif, mono)[0]
                    size = s["size"]
                    color = tuple(((s["color"] >> k) & 255) / 255 for k in (16, 8, 0))
                    o = fitz.Point(s["origin"])
                    tw = fitz.TextWriter(page.rect, color=color)
                    _append_words(tw, (o.x, o.y), text, font, size)
                    want = math.hypot(s["bbox"][2] - s["bbox"][0], s["bbox"][3] - s["bbox"][1]) if angle else s["bbox"][2] - s["bbox"][0]
                    got = font.text_length(text.rstrip(), fontsize=size)
                    sx = max(0.75, min(1.3, want / got)) if got > 0 and want > 0 and not angle else 1.0
                    m = fitz.Matrix(sx, 0, 0, 1, 0, 0)
                    if angle:
                        m = m * fitz.Matrix(angle)
                    if abs(sx - 1) < 0.01 and not angle:
                        tw.write_text(page)
                    else:
                        tw.write_text(page, morph=(o, m))
            page.clean_contents()                  # the fonts no longer used leave the resources
            if rot:
                page.set_rotation(rot)
            changed.append(i)
        return changed

    def save_pdfa(self, path):
        """A PDF/A-2b copy for long-term archiving (e.g. for public offices): fonts
        embedded, sRGB colour profile, metadata declared, no password, JavaScript or
        attachments. The open document is not changed."""
        with self.lock:
            done = []
            d = fitz.open(stream=self.doc.tobytes(encryption=fitz.PDF_ENCRYPT_NONE), filetype="pdf")
            if self._has_js(d) or d.embfile_count():
                done.append("Tolti JavaScript e allegati")
            d.scrub(attached_files=True, clean_pages=False, embedded_files=True, hidden_text=False, javascript=True,
                    metadata=False, redactions=False, remove_links=False, reset_fields=False, reset_responses=False,
                    thumbnails=False, xml_metadata=True)
            fixed = self._embed_missing_fonts(d)
            if fixed:
                done.append(f"Caratteri incorporati in {len(fixed)} pagine")
            for p in d:
                for a in p.annots() or []:
                    try:
                        a.set_flags(fitz.PDF_ANNOT_IS_PRINT)
                        a.update()
                    except Exception:
                        pass
            cat = d.pdf_catalog()
            af = d.xref_get_key(cat, "AcroForm")
            if af[0] == "xref":
                d.xref_set_key(int(af[1].split()[0]), "NeedAppearances", "null")
            elif af[0] == "dict":
                d.xref_set_key(cat, "AcroForm/NeedAppearances", "null")
            # sRGB output intent
            from PIL import ImageCms
            icc = ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB")).tobytes()
            ix = d.get_new_xref()
            d.update_object(ix, "<< /N 3 >>")
            d.update_stream(ix, icc)
            d.xref_set_key(cat, "OutputIntents", "[<< /Type /OutputIntent /S /GTS_PDFA1 /OutputConditionIdentifier "
                           f"({PDFA_RGB}) /Info ({PDFA_RGB}) /RegistryName (http://www.color.org) /DestOutputProfile {ix} 0 R >>]")
            # metadata, the same in the info dictionary and in XMP
            now = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0)
            m = d.metadata or {}
            title = (m.get("title") or "").strip() or os.path.splitext(os.path.basename(path))[0]
            info = {"title": title, "author": m.get("author") or "", "subject": m.get("subject") or "",
                    "keywords": m.get("keywords") or "", "creator": m.get("creator") or "Calamo PDF Studio",
                    "producer": "Calamo PDF Studio", "creationDate": _pdf_date(now), "modDate": _pdf_date(now)}
            d.set_metadata(info)
            iso = now.strftime("%Y-%m-%dT%H:%M:%S+00:00")
            e = _xmp_escape
            xmp = f"""<?xpacket begin="﻿" id="W5M0MpCehiHzreSzNTczkc9d"?>
<x:xmpmeta xmlns:x="adobe:ns:meta/">
 <rdf:RDF xmlns:rdf="http://www.w3.org/1999/02/22-rdf-syntax-ns#">
  <rdf:Description rdf:about="" xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:xmp="http://ns.adobe.com/xap/1.0/"
    xmlns:pdf="http://ns.adobe.com/pdf/1.3/" xmlns:pdfaid="http://www.aiim.org/pdfa/ns/id/">
   <dc:format>application/pdf</dc:format>
   <dc:title><rdf:Alt><rdf:li xml:lang="x-default">{e(title)}</rdf:li></rdf:Alt></dc:title>
   {f'<dc:creator><rdf:Seq><rdf:li>{e(info["author"])}</rdf:li></rdf:Seq></dc:creator>' if info["author"] else ''}
   {f'<dc:description><rdf:Alt><rdf:li xml:lang="x-default">{e(info["subject"])}</rdf:li></rdf:Alt></dc:description>' if info["subject"] else ''}
   {f'<pdf:Keywords>{e(info["keywords"])}</pdf:Keywords>' if info["keywords"] else ''}
   <pdf:Producer>{e(info["producer"])}</pdf:Producer>
   <xmp:CreatorTool>{e(info["creator"])}</xmp:CreatorTool>
   <xmp:CreateDate>{iso}</xmp:CreateDate>
   <xmp:ModifyDate>{iso}</xmp:ModifyDate>
   <xmp:MetadataDate>{iso}</xmp:MetadataDate>
   <pdfaid:part>2</pdfaid:part>
   <pdfaid:conformance>B</pdfaid:conformance>
  </rdf:Description>
 </rdf:RDF>
</x:xmpmeta>
<?xpacket end="w"?>"""
            d.set_xml_metadata(xmp)
            done.append("Profilo colore sRGB e metadati PDF/A-2b")
            d.save(path, garbage=3, deflate=True, encryption=fitz.PDF_ENCRYPT_NONE)
            check = self.preflight(fitz.open(path))
            left = [it for it in check["items"] if it["group"] == "pdfa" and not it["ok"]]
            return {"path": path, "done": done, "left": left, "ok": not left}

    # ================================================================== export
    def _ink_text(self, i):
        return "\n".join(l["text"] for l in self.ink.get(i, []) if l.get("text") and l.get("source") != "unread")

    def export_text(self, path):
        with self.lock:
            parts = []
            for i, page in enumerate(self.doc):
                t = page.get_text("text", sort=True).strip()
                ink = self._ink_text(i)
                if ink and len(t) < 20:
                    t = ink
                parts.append(f"— Pagina {i + 1} —\n{t}")
            with open(path, "w", encoding="utf-8") as f:
                f.write("\n\n".join(parts) + "\n")
            return {"path": path, "name": os.path.basename(path)}

    def export_html(self, path):
        with self.lock:
            title = (self.doc.metadata or {}).get("title") or os.path.splitext(os.path.basename(path))[0]
            body = []
            for i, page in enumerate(self.doc):
                body.append(f'<section class="pdf-page" aria-label="Pagina {i + 1}">{page.get_text("html")}</section>')
            html = f"""<!doctype html>
<html lang="{self._lang() or 'it'}"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{_xmp_escape(title)}</title>
<style>body{{background:#ecebe8;margin:0;padding:24px 0;font-family:sans-serif}}
.pdf-page{{width:max-content;max-width:100%;margin:0 auto 24px;background:#fff;box-shadow:0 2px 12px rgba(0,0,0,.18);overflow:auto}}
.pdf-page>div{{margin:0 auto}}</style></head>
<body>
{chr(10).join(body)}
</body></html>"""
            with open(path, "w", encoding="utf-8") as f:
                f.write(html)
            return {"path": path, "name": os.path.basename(path)}

    def export_pptx(self, path, progress=None):
        """One slide per page: the drawings and pictures of the page as the background,
        every line of text in an editable text box with its font, size and colour."""
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import MSO_AUTO_SIZE
        from pptx.util import Emu, Pt
        from .document import REDACT_KW
        with self.lock:
            EMU = 12700                                   # per point
            MAXPT = 4000
            first = self.doc[0].rect
            k0 = min(1.0, MAXPT / max(first.width, first.height))
            SW, SH = first.width * k0, first.height * k0
            prs = Presentation()
            prs.slide_width, prs.slide_height = Emu(int(SW * EMU)), Emu(int(SH * EMU))
            blank = prs.slide_layouts[6]
            n = len(self.doc)
            for i, page in enumerate(self.doc):
                if progress:
                    progress(i, n)
                R = page.rect
                k = min(SW / R.width, SH / R.height)
                ox, oy = (SW - R.width * k) / 2, (SH - R.height * k) / 2
                slide = prs.slides.add_slide(blank)
                # the background: the page without its text
                one = fitz.open()
                one.insert_pdf(self.doc, from_page=i, to_page=i, annots=True)
                p1 = one[0]
                d = page.get_text("dict")
                lines = [l for b in d["blocks"] if b.get("type") == 0 for l in b["lines"]
                         if "".join(s["text"] for s in l["spans"]).strip()]
                for l in lines:
                    p1.add_redact_annot(fitz.Rect(l["bbox"]))
                if lines:
                    p1.apply_redactions(**REDACT_KW)
                pix = p1.get_pixmap(dpi=150, alpha=False, annots=True)
                arr = np.frombuffer(pix.samples, dtype=np.uint8)
                if arr.size and arr.min() < 245:
                    import io
                    slide.shapes.add_picture(io.BytesIO(pix.tobytes("png")), Emu(int(ox * EMU)), Emu(int(oy * EMU)),
                                             Emu(int(R.width * k * EMU)), Emu(int(R.height * k * EMU)))
                for l in lines:
                    bb = self._shown(page, fitz.Rect(l["bbox"]))
                    x, y = ox + bb.x0 * k, oy + bb.y0 * k
                    w, h = max(bb.width * k * 1.08, 6), max(bb.height * k, 4)
                    tb = slide.shapes.add_textbox(Emu(int(x * EMU)), Emu(int(y * EMU)), Emu(int(w * EMU)), Emu(int(h * EMU)))
                    tf = tb.text_frame
                    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                    tf.word_wrap = False
                    tf.auto_size = MSO_AUTO_SIZE.NONE      # the box stays where the line was
                    para = tf.paragraphs[0]
                    for s in l["spans"]:
                        if not s["text"]:
                            continue
                        run = para.add_run()
                        run.text = s["text"]
                        f = run.font
                        f.size = Pt(max(1, round(s["size"] * k * 2) / 2))
                        bold, italic, _, _ = F.style_from(F.clean_name(s["font"]), s["flags"])
                        f.bold, f.italic = bold, italic
                        f.name = F.css_family(s["font"], s["flags"])
                        c = s["color"]
                        f.color.rgb = RGBColor((c >> 16) & 255, (c >> 8) & 255, c & 255)
            prs.save(path)
            return {"path": path, "name": os.path.basename(path), "slides": n}

    def remove_password(self):
        """The open document is saved without its password from now on."""
        with self.lock:
            if not getattr(self, "_password", None):
                raise RuntimeError("Questo documento non ha una password.")
            self._password = None
            self.dirty = True
            self.version += 1
            return self.info()

    # ================================================================== digital signatures
    def _signed_bytes(self):
        """The bytes the signatures refer to: the file on disk when nothing changed (it must
        not be rewritten, or the signatures already in it would break)."""
        if self.path and not self.dirty and os.path.exists(self.path) and not getattr(self, "_password", None):
            with open(self.path, "rb") as f:
                return f.read()
        return self.doc.tobytes(garbage=1, deflate=True, encryption=fitz.PDF_ENCRYPT_NONE)

    def sign_digital(self, out_path, cert_path, password, index=0, rect=None, field=None, reason="",
                     location="", tsa_url=""):
        """Sign with a certificate into a new file (rect in displayed coordinates)."""
        from . import signing
        with self.lock:
            box = None
            if not field:
                if not rect:
                    raise RuntimeError("Trascina sulla pagina il riquadro della firma.")
                page = self.doc[int(index)]
                r = fitz.Rect(rect)
                if r.width < 40 or r.height < 18:
                    raise RuntimeError("Il riquadro della firma è troppo piccolo.")
                # displayed -> PDF user space (origin bottom left)
                u = (r * page.derotation_matrix) * ~page.transformation_matrix
                u.normalize()
                box = (u.x0, u.y0, u.x1, u.y1)
            data = self._signed_bytes()
        return signing.sign(data, out_path, cert_path, password, page=int(index), box=box, field=field,
                            reason=reason, location=location, tsa_url=tsa_url)

    def verify_signatures(self):
        from . import signing
        with self.lock:
            data = self._signed_bytes()
            if b"/ByteRange" not in data:
                return {"signatures": [], "dirty": self.dirty}
        return {"signatures": signing.verify(data), "dirty": bool(self.dirty and self.path)}
